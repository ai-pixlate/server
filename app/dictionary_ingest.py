"""규제사전·현지부적합 사전(xlsx) 검증·적재 CLI.

    python -m app.dictionary_ingest validate --regulatory regulation_dict.xlsx --local locale_unsuitable_dict.xlsx
    python -m app.dictionary_ingest plan     ...   # DB와 비교만(읽기 전용, 명시적 쓰기 차단 잠금 없음)
    python -m app.dictionary_ingest load     ...   # 잠금 후 비교를 다시 하고 사전·근거를 한 트랜잭션으로 적재

파일은 명시한 경로만 읽는다(하나 이상). 원본은 수정하지 않는다. DB 접속은 DATABASE_URL(plan·load).

읽는 탭(데이터팀 00_읽는법): 규제사전 파일은 「규제사전」·「근거」, 현지부적합 파일은 「01_현지부적합사전」.

규제사전 → expression_dictionary (dict_type=regulatory)
- 필수(●) 열이 비면 오류. 시트 전용 열은 저장하지 않지만 규칙은 검사한다:
  confidence 는 high 만, version 은 양의 정수, internal_category 는 9월엔 비어 있어야 한다.
- variant_expressions.ko / .en 은 '; ' 로 나눠 variant_ko / forbidden_en JSON 배열로 저장.
- verdict_status rewritable 은 regulated 로 바꿔 저장(완충형, 9/21 PM 결정기록). 대체 표현이 비면 오류.
  시트의 판정 원값은 source_verdict_status 에 그대로 남긴다(서비스 판정은 verdict_status). 마지막 적재 원값이며
  변경 이력이 아니다.
  allowed 는 대체 표현이 비어야 하고, conditional·rewritable 은 채워져야 한다.
- 본문 verified_at → confirmed_date.
- 대표 근거: 행의 evidence_id 가 「근거」 탭에 있고, is_primary=Y 이며 rg_id 에 이 항목이 들어 있어야 하고,
  source_type·article·url 이 행의 evidence_* 와 같아야 한다. 이 항목을 가리키는 is_primary=Y 는 정확히 1건.
  근거 탭의 필수값(evidence_id·source_type·quote·url·verified_at)은 적재할 대표 근거 행에서 검사한다.
  근거 verified_at 은 검사만 하고 저장하지 않는다(근거 테이블에 열이 없음). 시장 관행은 allowed 행에만.
- 대체 표현에 한국어가 남은 행은 적재하되 '미해결 — 프롬프트 사용 불가'로 따로 보고한다.

현지부적합 → expression_dictionary (dict_type=local, target_country=US, regulatory_class=common)
- common = 미국 작업의 모든 규제 분류에 공통 적용. 국가 범위를 없애는 뜻이 아니다.
- 항목→source_expression · 패턴→variant_ko · 판정→verdict_status(원값도 source_verdict_status) · 제외하는/제외하지 않는 맥락 →
  exclusion_context/keep_context · 셀러 문장→reason · verified_at→confirmed_date. 근거 행은 없다.
- 판단 근거·kr_freq·kr_corpus 는 이번 DB 적재에서 제외하고 원본 파일로 보존한다
  (DB만으로 판정의 내부 배경을 복원할 수 없다).

재적재(같은 external_id):
- 사전: 내부 PK 유지. 모든 열이 같으면 변경 없음, 다르면 갱신. 같은 ID의 dict_type·국가·규제 분류가
  바뀌면 오류(번호 재사용 의심). 같은 내용 키(dict_type, 국가, 분류, source_expression)에 다른 ID면 오류.
- 대표 근거: 사전 항목마다 (dictionary_id, 근거 ID) 연결 행의 PK 유지. 대표가 다른 근거로 바뀌면
  이전 대표 연결은 삭제하지 않고 비대표(is_primary=false)로 남긴다. 근거 내용의 변경 이력은 관리하지 않는다
  (같은 연결의 내용이 바뀌면 그 행을 갱신한다). 적재 후 규제 항목마다 대표 근거가 정확히 1건인지 검사.
- 파일에 없는 기존 행은 지우지 않고 건수만 보고한다.
- 제한: DB에 version 을 저장하지 않으므로 옛 파일로 최신 내용을 덮어쓰는 것을 막지 못한다.
  출력되는 파일 SHA-256 은 파일 식별용이지 역행 방지 장치가 아니다. 적재한 원본 파일과 실행 기록을 보관한다.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from typing import Any, Optional

import openpyxl
from sqlalchemy import create_engine, text
from sqlalchemy.engine import Connection
from sqlalchemy.exc import DBAPIError

from app.glossary_ingest import CellError, Issue, _blank, _date, _text, print_issues

REG_SHEET = "규제사전"
EVIDENCE_SHEET = "근거"
LOCAL_SHEET = "01_현지부적합사전"

REG_COLUMNS = [
    "id", "dict_type", "target_country", "regulatory_class", "source_expression",
    "variant_expressions.ko", "variant_expressions.en", "alternative_expression", "verdict_status", "reason",
    "evidence_id", "evidence_source_type", "evidence_article", "evidence_url", "verified_at",
    "internal_category", "confidence", "source", "version",
]
EVIDENCE_COLUMNS = ["evidence_id", "rg_id", "source_type", "company", "issued_at", "quote", "article", "url",
                    "is_primary", "verified_at"]
LOCAL_COLUMNS = ["id", "항목", "패턴", "판정", "제외하는 맥락", "제외하지 않는 맥락", "셀러 문장", "verified_at"]

DICT_CLASSES = ("cosmetic", "otc", "common")  # 계약 정본 §4-1: 사전 행의 적용 분류
REG_VERDICTS = ("regulated", "conditional", "rewritable", "allowed")  # 시트 값. rewritable → regulated
LOCAL_VERDICTS = ("irrelevant", "needs_fix", "cultural")
# 원천 탭별 근거 ID 접두어와 source_type (00_읽는법 '원천 탭 공통 규칙')
EVIDENCE_TYPES = {"WL": "Warning Letter", "GD": "FDA 공식 문서", "MN": "OTC 모노그래프", "MP": "시장 관행"}
MARKET_PRACTICE = "시장 관행"
LOCAL_COUNTRY, LOCAL_CLASS = "US", "common"

DICT_FIELDS = [
    "external_id", "dict_type", "target_country", "regulatory_class", "source_expression", "variant_ko",
    "forbidden_en", "verdict_status", "alternative_expression", "reason", "confirmed_date",
    "exclusion_context", "keep_context", "source_verdict_status",
]
SCOPE_KEY = ("dict_type", "target_country", "regulatory_class")
CONTENT_KEY = SCOPE_KEY + ("source_expression",)
EVIDENCE_FIELDS = ["evidence_source_type", "evidence_document", "evidence_quote", "evidence_article",
                   "evidence_url", "is_primary"]
VARCHAR_LIMITS = {"external_id": 32, "dict_type": 20, "target_country": 8, "regulatory_class": 30,
                  "source_expression": 300, "verdict_status": 20, "source_verdict_status": 20}
HANGUL = re.compile(r"[가-힣]")


@dataclass
class DictFile:
    kind: str  # regulatory | local
    path: str
    sha256: str = ""
    total_rows: int = 0
    records: list[dict[str, Any]] = field(default_factory=list)
    rows: list[int] = field(default_factory=list)
    evidence: dict[str, dict[str, Any]] = field(default_factory=dict)  # external_id(RG) → 대표 근거
    converted: list[str] = field(default_factory=list)  # rewritable → regulated
    unresolved: list[str] = field(default_factory=list)  # 대체 표현에 한국어가 남은 항목
    versions: Counter = field(default_factory=Counter)
    evidence_skipped: int = 0  # 「규제사전」 탭에 없는 RG(보류·백업 등)만 가리키는 근거 행


@dataclass
class DictValidation:
    files: list[DictFile]
    issues: list[Issue]

    @property
    def records(self) -> list[dict[str, Any]]:
        return [r for f in self.files for r in f.records]


def _sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fp:
        for chunk in iter(lambda: fp.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def _split(v: Any, required: bool) -> Optional[list[str]]:
    s = _text(v, required)
    if s is None:
        return None
    parts = s.split("; ")
    if any(p.strip() == "" or ";" in p for p in parts):
        raise CellError("'; '(세미콜론+한 칸)로 나눈 값에 빈 값이나 다른 구분자가 있음")
    return parts


def _version(v: Any) -> int:
    if _blank(v):
        raise CellError("필수값이 비어 있음")
    if isinstance(v, bool):
        raise CellError(f"양의 정수가 아님({v!r})")
    if isinstance(v, int):
        n = v
    elif isinstance(v, float) and v.is_integer():
        n = int(v)
    elif isinstance(v, str) and v.strip().isdigit():
        n = int(v.strip())
    else:
        raise CellError(f"양의 정수가 아님({v!r})")
    if n < 1:
        raise CellError(f"양의 정수가 아님({v!r})")
    return n


def _required_date(v: Any) -> dt.date:
    d = _date(v)
    if d is None:
        raise CellError("필수값이 비어 있음")
    return d


def _read_sheet(wb, sheet: str, wanted: list[str], label: str, issues: list[Issue]):
    """(열 → 위치, [(엑셀 행, 셀 dict)]) 또는 열 오류 시 None."""
    if sheet not in wb.sheetnames:
        issues.append(Issue(label, None, None, None, f"시트 '{sheet}' 없음"))
        return None
    it = wb[sheet].iter_rows(values_only=True)
    header = [None if h is None else str(h).strip() for h in next(it, ())]
    idx: dict[str, int] = {}
    ok = True
    for col in wanted:
        pos = [i for i, h in enumerate(header) if h == col]
        if len(pos) != 1:
            issues.append(Issue(label, 1, None, col, "필수 열 없음" if not pos else "같은 이름의 열이 여러 개"))
            ok = False
        else:
            idx[col] = pos[0]
    if not ok:
        return None
    out = []
    for row_no, row in enumerate(it, start=2):
        if all(_blank(v) for v in row):
            continue
        out.append((row_no, {c: (row[i] if i < len(row) else None) for c, i in idx.items()}))
    return out


class _Row:
    """한 행의 값 변환과 오류 모음."""

    def __init__(self, label, row_no, eid, issues):
        self.label, self.row_no, self.eid = label, row_no, eid
        self.errs: list[Issue] = []
        self.issues = issues

    def err(self, col, msg):
        self.errs.append(Issue(self.label, self.row_no, self.eid, col, msg))

    def take(self, rec, cell, col, key, fn):
        try:
            rec[key] = fn(cell[col])
        except CellError as e:
            self.err(col, str(e))

    def finish(self) -> bool:
        self.issues.extend(self.errs)
        return not self.errs


def _limits(r: _Row, rec: dict) -> None:
    for k, limit in VARCHAR_LIMITS.items():
        v = rec.get(k)
        if isinstance(v, str) and len(v) > limit:
            r.err(k, f"{limit}자 초과({len(v)}자)")


def parse_regulatory(path: str, issues: list[Issue]) -> DictFile:
    df = DictFile("regulatory", path)
    label = f"regulatory({os.path.basename(path)})"
    try:
        df.sha256 = _sha256(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        issues.append(Issue(label, None, None, None, f"파일을 열 수 없음: {e}"))
        return df
    try:
        reg = _read_sheet(wb, REG_SHEET, REG_COLUMNS, f"{label}/{REG_SHEET}", issues)
        ev = _read_sheet(wb, EVIDENCE_SHEET, EVIDENCE_COLUMNS, f"{label}/{EVIDENCE_SHEET}", issues)
    finally:
        wb.close()
    if reg is None or ev is None:
        return df

    # 근거 탭: ID 형식·중복은 모든 행에서 검사
    ev_label = f"{label}/{EVIDENCE_SHEET}"
    ev_rows: dict[str, tuple[int, dict]] = {}
    rg_refs: dict[str, list[str]] = defaultdict(list)  # RG → 그 RG를 가리키는 is_primary=Y 근거 ID
    for row_no, cell in ev:
        eid = cell["evidence_id"]
        m = re.match(r"^(WL|GD|MN|MP)-\d+$", eid) if isinstance(eid, str) else None
        if not m:
            issues.append(Issue(ev_label, row_no, None, "evidence_id", f"WL/GD/MN/MP-숫자 형식이 아님({eid!r})"))
            continue
        if eid in ev_rows:
            issues.append(Issue(ev_label, row_no, eid, "evidence_id", f"근거 ID 중복(첫 행 {ev_rows[eid][0]})"))
            continue
        ev_rows[eid] = (row_no, cell)
        if cell["is_primary"] not in (None, "Y", "N"):
            issues.append(Issue(ev_label, row_no, eid, "is_primary", f"Y/N 이 아님({cell['is_primary']!r})"))
        rg = cell["rg_id"]
        if isinstance(rg, str) and rg.strip() and cell["is_primary"] == "Y":
            for x in rg.split("; "):
                rg_refs[x.strip()].append(eid)

    reg_label = f"{label}/{REG_SHEET}"
    sheet_ids = {c["id"] for _, c in reg if isinstance(c["id"], str)}
    ev_checked: dict[str, tuple[Optional[dict], int]] = {}  # 근거 ID → (저장값, 오류 수) — 공유 근거는 한 번만 검사
    for row_no, cell in reg:
        df.total_rows += 1
        raw = cell["id"]
        eid = raw if isinstance(raw, str) else None
        r = _Row(reg_label, row_no, eid, issues)
        if eid is None or not re.match(r"^RG-\d+$", eid):
            r.err("id", f"RG-숫자 형식이 아님({raw!r})")
        rec: dict[str, Any] = {"external_id": eid, "exclusion_context": None, "keep_context": None}
        r.take(rec, cell, "dict_type", "dict_type", lambda v: _text(v, True))
        if rec.get("dict_type") not in (None, "regulatory"):
            r.err("dict_type", f"이 탭은 regulatory 만 허용({rec['dict_type']!r})")
        r.take(rec, cell, "target_country", "target_country", lambda v: _text(v, True))
        r.take(rec, cell, "regulatory_class", "regulatory_class", lambda v: _text(v, True))
        if rec.get("regulatory_class") not in (None, *DICT_CLASSES):
            r.err("regulatory_class", f"{'·'.join(DICT_CLASSES)} 가 아님({rec['regulatory_class']!r})")
        r.take(rec, cell, "source_expression", "source_expression", lambda v: _text(v, True))
        r.take(rec, cell, "variant_expressions.ko", "variant_ko", lambda v: _split(v, True))
        r.take(rec, cell, "variant_expressions.en", "forbidden_en", lambda v: _split(v, True))
        r.take(rec, cell, "alternative_expression", "alternative_expression", lambda v: _text(v, False))
        r.take(rec, cell, "reason", "reason", lambda v: _text(v, True))
        r.take(rec, cell, "verified_at", "confirmed_date", _required_date)
        try:
            ver = _version(cell["version"])
        except CellError as e:
            r.err("version", str(e))
            ver = None
        if cell["confidence"] != "high":
            r.err("confidence", f"9월 시드는 high 만 적재({cell['confidence']!r})")
        if not _blank(cell["internal_category"]):
            r.err("internal_category", "9월에는 비워 둔다(값 체계 미확정)")

        status = cell["verdict_status"]
        alt = rec.get("alternative_expression")
        if status not in REG_VERDICTS:
            r.err("verdict_status", f"{'·'.join(REG_VERDICTS)} 가 아님({status!r})")
        else:
            if status == "allowed" and alt is not None:
                r.err("alternative_expression", "allowed 는 대체 표현을 비운다")
            if status in ("rewritable", "conditional") and alt is None:
                r.err("alternative_expression", f"{status} 는 대체 표현이 있어야 한다")
            rec["verdict_status"] = "regulated" if status == "rewritable" else status
            rec["source_verdict_status"] = status

        # 대표 근거
        ev_id = cell["evidence_id"]
        evidence = None
        if not isinstance(ev_id, str) or ev_id not in ev_rows:
            r.err("evidence_id", f"「근거」 탭에 없는 근거 ID({ev_id!r})")
        else:
            ev_row, ec = ev_rows[ev_id]
            prim = rg_refs.get(eid or "", [])
            if prim != [ev_id]:
                r.err("evidence_id", f"이 항목을 가리키는 is_primary=Y 근거가 {prim} — {ev_id} 하나여야 함")
            for dcol, ecol in (("evidence_source_type", "source_type"), ("evidence_article", "article"),
                               ("evidence_url", "url")):
                if cell[dcol] != ec[ecol]:
                    r.err(dcol, f"근거 {ev_id}({ev_row}행)의 {ecol} 와 다름")
            prefix_type = EVIDENCE_TYPES[ev_id[:2]]
            if ec["source_type"] != prefix_type:
                r.err("evidence_source_type", f"{ev_id} 는 {prefix_type} 이어야 함({ec['source_type']!r})")
            if ec["source_type"] == MARKET_PRACTICE and status != "allowed":
                r.err("evidence_source_type", "시장 관행 근거는 allowed 행에만 쓸 수 있다")
            if ev_id not in ev_checked:
                er = _Row(ev_label, ev_row, ev_id, issues)
                got = {"external_id": ev_id, "evidence_document": None, "is_primary": True}
                er.take(got, ec, "source_type", "evidence_source_type", lambda v: _text(v, True))
                er.take(got, ec, "quote", "evidence_quote", lambda v: _text(v, True))
                er.take(got, ec, "article", "evidence_article", lambda v: _text(v, False))
                er.take(got, ec, "url", "evidence_url", lambda v: _text(v, True))
                er.take({}, ec, "verified_at", "verified_at", _required_date)  # 검사만, 저장 안 함
                ev_checked[ev_id] = (got, len(er.errs))
                er.finish()
            evidence, n_err = ev_checked[ev_id]
            if n_err:
                r.err("evidence_id", f"대표 근거 {ev_id} 에 필수값 오류 {n_err}건(「근거」 탭 오류 참조)")

        _limits(r, rec)
        if r.finish():
            df.records.append(rec)
            df.rows.append(row_no)
            df.evidence[eid] = evidence
            df.versions[ver] += 1
            if status == "rewritable":
                df.converted.append(eid)
            if alt and HANGUL.search(alt):
                df.unresolved.append(eid)

    df.evidence_skipped = sum(
        1 for _, c in ev_rows.values()
        if isinstance(c["rg_id"], str) and c["rg_id"].strip()
        and not any(x.strip() in sheet_ids for x in c["rg_id"].split("; "))
    )
    return df


def parse_local(path: str, issues: list[Issue]) -> DictFile:
    df = DictFile("local", path)
    label = f"local({os.path.basename(path)})"
    try:
        df.sha256 = _sha256(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        issues.append(Issue(label, None, None, None, f"파일을 열 수 없음: {e}"))
        return df
    try:
        rows = _read_sheet(wb, LOCAL_SHEET, LOCAL_COLUMNS, f"{label}/{LOCAL_SHEET}", issues)
    finally:
        wb.close()
    if rows is None:
        return df
    lbl = f"{label}/{LOCAL_SHEET}"
    for row_no, cell in rows:
        df.total_rows += 1
        raw = cell["id"]
        eid = raw if isinstance(raw, str) else None
        r = _Row(lbl, row_no, eid, issues)
        if eid is None or not re.match(r"^LC-\d+$", eid):
            r.err("id", f"LC-숫자 형식이 아님({raw!r})")
        rec: dict[str, Any] = {
            "external_id": eid, "dict_type": "local", "target_country": LOCAL_COUNTRY,
            "regulatory_class": LOCAL_CLASS, "forbidden_en": None, "alternative_expression": None,
        }
        r.take(rec, cell, "항목", "source_expression", lambda v: _text(v, True))
        r.take(rec, cell, "패턴", "variant_ko", lambda v: _split(v, True))
        r.take(rec, cell, "판정", "verdict_status", lambda v: _text(v, True))
        if rec.get("verdict_status") not in (None, *LOCAL_VERDICTS):
            r.err("판정", f"{'·'.join(LOCAL_VERDICTS)} 가 아님({rec['verdict_status']!r})")
        rec["source_verdict_status"] = rec.get("verdict_status")
        r.take(rec, cell, "제외하는 맥락", "exclusion_context", lambda v: _text(v, True))
        r.take(rec, cell, "제외하지 않는 맥락", "keep_context", lambda v: _text(v, True))
        r.take(rec, cell, "셀러 문장", "reason", lambda v: _text(v, True))
        r.take(rec, cell, "verified_at", "confirmed_date", _required_date)
        _limits(r, rec)
        if r.finish():
            df.records.append(rec)
            df.rows.append(row_no)
    return df


def validate_files(regulatory: Optional[str], local: Optional[str]) -> DictValidation:
    issues: list[Issue] = []
    files = []
    if regulatory:
        files.append(parse_regulatory(regulatory, issues))
    if local:
        files.append(parse_local(local, issues))
    by_id: dict[str, list] = defaultdict(list)
    by_key: dict[tuple, list] = defaultdict(list)
    for f in files:
        for rec, row in zip(f.records, f.rows):
            by_id[rec["external_id"]].append((f.kind, row))
            by_key[tuple(rec[k] for k in CONTENT_KEY)].append((f.kind, row, rec["external_id"]))
    for eid, where in by_id.items():
        if len(where) > 1:
            issues.append(Issue("전체", None, eid, "id", f"항목 ID 중복: {where}"))
    for key, where in by_key.items():
        if len(where) > 1:
            issues.append(Issue("전체", None, None, None, f"내용 키 중복 {key}: {where}"))
    return DictValidation(files, issues)


# ── DB 비교·적재 ──────────────────────────────────────────────────────────

@dataclass
class DictPlan:
    inserts: list[dict] = field(default_factory=list)
    updates: list[dict] = field(default_factory=list)
    unchanged: int = 0
    db_only: int = 0
    ev_insert: list[dict] = field(default_factory=list)  # dictionary external_id 로 연결(적재 후 id 해석)
    ev_update: list[dict] = field(default_factory=list)
    ev_demote: list[int] = field(default_factory=list)  # 근거 행 id
    ev_unchanged: int = 0
    issues: list[Issue] = field(default_factory=list)


def _json_list(v):
    return None if v is None else json.dumps(v, ensure_ascii=False)


def _db_dict_rows(conn: Connection):
    return conn.execute(text(f"SELECT id, {', '.join(DICT_FIELDS)} FROM expression_dictionary")).mappings().all()


def _require_schema(conn: Connection) -> None:
    cols = {(r[0], r[1]) for r in conn.execute(text(
        "SELECT table_name, column_name FROM information_schema.columns WHERE table_schema = current_schema() "
        "AND table_name IN ('expression_dictionary', 'expression_dictionary_evidence')"))}
    need = [("expression_dictionary", c) for c in DICT_FIELDS] + [("expression_dictionary_evidence", "external_id")]
    missing = [f"{t}.{c}" for t, c in need if (t, c) not in cols]
    if missing:
        raise SystemExit(f"열이 없음 {missing} — 마이그레이션 0006 을 먼저 적용")


def build_plan(conn: Connection, vr: DictValidation) -> DictPlan:
    plan = DictPlan()
    existing = _db_dict_rows(conn)
    by_ext = {r["external_id"]: r for r in existing}
    by_key = {tuple(r[k] for k in CONTENT_KEY): r for r in existing}
    ev_rows = conn.execute(text(
        "SELECT id, dictionary_id, external_id, " + ", ".join(EVIDENCE_FIELDS) + " FROM expression_dictionary_evidence"
    )).mappings().all()
    ev_by_dict: dict[int, list] = defaultdict(list)
    for e in ev_rows:
        ev_by_dict[e["dictionary_id"]].append(e)

    incoming = set()
    for f in vr.files:
        for rec, row in zip(f.records, f.rows):
            eid = rec["external_id"]
            incoming.add(eid)

            def fail(msg):
                plan.issues.append(Issue(f.kind, row, eid, None, msg))

            cur = by_ext.get(eid)
            if cur is not None:
                if tuple(cur[k] for k in SCOPE_KEY) != tuple(rec[k] for k in SCOPE_KEY):
                    fail(f"같은 항목 ID인데 dict_type·국가·규제 분류가 바뀜(번호 재사용 의심) DB id={cur['id']}")
                    continue
                other = by_key.get(tuple(rec[k] for k in CONTENT_KEY))
                if other is not None and other["id"] != cur["id"]:
                    fail(f"같은 내용 키에 다른 항목 ID가 이미 있음({other['external_id']})")
                    continue
                if all(cur[k] == rec[k] for k in DICT_FIELDS):
                    plan.unchanged += 1
                else:
                    plan.updates.append({**rec, "id": cur["id"]})
            else:
                other = by_key.get(tuple(rec[k] for k in CONTENT_KEY))
                if other is not None:
                    fail(f"같은 내용 키에 다른 항목 ID가 이미 있음({other['external_id']})")
                    continue
                plan.inserts.append(rec)

            ev = f.evidence.get(eid)
            if ev is None:
                continue
            rows = ev_by_dict.get(cur["id"], []) if cur is not None else []
            if any(e["is_primary"] and e["external_id"] is None for e in rows):
                fail("원본 근거 ID 없는 기존 대표 근거가 있음 — 연결 방법 결정 필요")
                continue
            for e in rows:
                if e["is_primary"] and e["external_id"] != ev["external_id"]:
                    plan.ev_demote.append(e["id"])
            match = next((e for e in rows if e["external_id"] == ev["external_id"]), None)
            want = {k: ev[k] for k in EVIDENCE_FIELDS}
            if match is None:
                plan.ev_insert.append({**want, "dict_external_id": eid, "external_id": ev["external_id"]})
            elif all(match[k] == want[k] for k in EVIDENCE_FIELDS):
                plan.ev_unchanged += 1
            else:
                plan.ev_update.append({**want, "id": match["id"]})

    kinds = {f.kind for f in vr.files}
    plan.db_only = sum(1 for r in existing if r["dict_type"] in kinds and r["external_id"] not in incoming)
    return plan


def _dict_params(rec: dict) -> dict:
    p = {k: rec[k] for k in DICT_FIELDS}
    p["variant_ko"] = _json_list(rec["variant_ko"])
    p["forbidden_en"] = _json_list(rec["forbidden_en"])
    return p


def apply_plan(conn: Connection, plan: DictPlan) -> None:
    cols = ", ".join(DICT_FIELDS)
    vals = ", ".join(f"CAST(:{c} AS jsonb)" if c in ("variant_ko", "forbidden_en") else f":{c}" for c in DICT_FIELDS)
    if plan.inserts:
        conn.execute(text(f"INSERT INTO expression_dictionary ({cols}) VALUES ({vals})"),
                     [_dict_params(r) for r in plan.inserts])
    if plan.updates:
        sets = ", ".join(
            f"{c} = CAST(:{c} AS jsonb)" if c in ("variant_ko", "forbidden_en") else f"{c} = :{c}"
            for c in DICT_FIELDS if c != "external_id")
        conn.execute(text(f"UPDATE expression_dictionary SET {sets} WHERE id = :id AND external_id = :external_id"),
                     [{**_dict_params(r), "id": r["id"]} for r in plan.updates])
    # 대표 해제를 먼저(사전 항목당 대표 1건 부분 UNIQUE 인덱스)
    if plan.ev_demote:
        conn.execute(text("UPDATE expression_dictionary_evidence SET is_primary = false WHERE id = ANY(:ids)"),
                     {"ids": plan.ev_demote})
    ev_sets = ", ".join(f"{c} = :{c}" for c in EVIDENCE_FIELDS)
    if plan.ev_update:
        conn.execute(text(f"UPDATE expression_dictionary_evidence SET {ev_sets} WHERE id = :id"), plan.ev_update)
    if plan.ev_insert:
        ids = dict(conn.execute(text(
            "SELECT external_id, id FROM expression_dictionary WHERE external_id = ANY(:e)"),
            {"e": [r["dict_external_id"] for r in plan.ev_insert]}).all())
        cols = ["dictionary_id", "external_id"] + EVIDENCE_FIELDS
        conn.execute(
            text(f"INSERT INTO expression_dictionary_evidence ({', '.join(cols)}) "
                 f"VALUES ({', '.join(':' + c for c in cols)})"),
            [{**{k: r[k] for k in EVIDENCE_FIELDS}, "external_id": r["external_id"],
              "dictionary_id": ids[r["dict_external_id"]]} for r in plan.ev_insert])


def verify_loaded(conn: Connection, vr: DictValidation) -> list[Issue]:
    """커밋 전: 이번 파일의 항목이 파일 값 그대로 있고, 규제 항목은 대표 근거가 정확히 1건(파일과 같은 것),
    현지 항목은 근거가 없는지 확인한다."""
    issues = []
    db = {r["external_id"]: r for r in _db_dict_rows(conn)}
    ev = defaultdict(list)
    for e in conn.execute(text(
            "SELECT dictionary_id, external_id, " + ", ".join(EVIDENCE_FIELDS) + " FROM expression_dictionary_evidence"
    )).mappings():
        ev[e["dictionary_id"]].append(e)
    for f in vr.files:
        for rec in f.records:
            eid = rec["external_id"]
            cur = db.get(eid)
            if cur is None or any(cur[k] != rec[k] for k in DICT_FIELDS):
                issues.append(Issue("검증", None, eid, None, "적재 후 사전 값이 파일과 다름"))
                continue
            prim = [e for e in ev[cur["id"]] if e["is_primary"]]
            want = f.evidence.get(eid)
            if want is None:
                if ev[cur["id"]]:
                    issues.append(Issue("검증", None, eid, None, "현지 항목에 근거 행이 있음"))
            elif len(prim) != 1 or prim[0]["external_id"] != want["external_id"] or any(
                    prim[0][k] != want[k] for k in EVIDENCE_FIELDS):
                issues.append(Issue("검증", None, eid, None, f"대표 근거가 {len(prim)}건이거나 파일과 다름"))
    return issues


def fingerprint(conn: Connection, vr: DictValidation) -> tuple[int, int, str]:
    """이번 파일 항목의 (항목 ID:사전 PK)와 대표 근거 (항목 ID/근거 ID:근거 PK)의 md5."""
    ids = [r["external_id"] for r in vr.records]
    d = conn.execute(text("SELECT external_id, id FROM expression_dictionary WHERE external_id = ANY(:i)"),
                     {"i": ids}).all()
    e = conn.execute(text(
        "SELECT d.external_id, v.external_id, v.id FROM expression_dictionary_evidence v "
        "JOIN expression_dictionary d ON d.id = v.dictionary_id WHERE d.external_id = ANY(:i) AND v.is_primary"),
        {"i": ids}).all()
    joined = ",".join(f"{a}:{b}" for a, b in sorted(d)) + "|" + ",".join(f"{a}/{b}:{c}" for a, b, c in sorted(e))
    return len(d), len(e), hashlib.md5(joined.encode("utf-8")).hexdigest()


# ── 출력 ─────────────────────────────────────────────────────────────────

def print_summary(vr: DictValidation) -> None:
    print("== 파일 검증 ==")
    for f in vr.files:
        print(f"- {f.kind}: {os.path.basename(f.path)} sha256 {f.sha256}")
        verdicts = Counter(r["verdict_status"] for r in f.records)
        classes = Counter(r["regulatory_class"] for r in f.records)
        print(f"    데이터 {f.total_rows}행 · 적재 대상 {len(f.records)}행 · 판정 {dict(verdicts)} · 분류 {dict(classes)}")
        if f.kind == "regulatory":
            print(f"    대표 근거 {len(f.evidence)}건(원본 근거 ID {len({e['external_id'] for e in f.evidence.values()})}종)"
                  f" · 「규제사전」 탭에 없는 항목만 가리키는 근거 행 {f.evidence_skipped}개는 읽지 않음")
            print(f"    rewritable→regulated {len(f.converted)}행: {', '.join(f.converted) or '없음'}"
                  f" · version 분포 {dict(sorted(f.versions.items()))}")
            if f.unresolved:
                print(f"    ⚠ 미해결 — 대체 표현에 한국어가 남음(프롬프트 사용 불가, 데이터 정정 필요): "
                      f"{', '.join(f.unresolved)}")


def main(argv: Optional[list[str]] = None, database_url: Optional[str] = None) -> int:
    p = argparse.ArgumentParser(prog="python -m app.dictionary_ingest", description="규제사전·현지부적합 xlsx 검증·적재")
    p.add_argument("mode", choices=["validate", "plan", "load"],
                   help="validate=파일만 검사(DB 없음) · plan=DB와 비교만(읽기 전용) · load=적재")
    p.add_argument("--regulatory", metavar="XLSX", help="「규제사전」·「근거」 탭이 있는 파일")
    p.add_argument("--local", metavar="XLSX", help="「01_현지부적합사전」 탭이 있는 파일")
    p.add_argument("--max-errors", type=int, default=50)
    args = p.parse_args(argv)
    if not (args.regulatory or args.local):
        p.error("--regulatory / --local 중 하나 이상 필요")

    vr = validate_files(args.regulatory, args.local)
    print_summary(vr)
    if vr.issues:
        print_issues(vr.issues, args.max_errors)
        print("결과: 실패 — 파일 오류가 있어 DB에 쓰지 않음")
        return 1
    if args.mode == "validate":
        print("결과: 파일 검증 통과 (DB 미접속)")
        return 0

    url = database_url or os.getenv("DATABASE_URL")
    if not url:
        print("DATABASE_URL 환경변수가 없음")
        return 2
    engine = create_engine(url, future=True)
    try:
        with engine.connect() as conn:
            trans = conn.begin()
            try:
                if args.mode == "plan":
                    conn.execute(text("SET TRANSACTION READ ONLY"))  # 명시적 쓰기 차단 잠금 없음
                _require_schema(conn)
                if args.mode == "load":
                    conn.execute(text("SET LOCAL lock_timeout = '10s'"))
                    conn.execute(text("LOCK TABLE expression_dictionary, expression_dictionary_evidence "
                                      "IN SHARE ROW EXCLUSIVE MODE"))
                plan = build_plan(conn, vr)
                print("== DB 비교 ==")
                print(f"- 사전: 삽입 {len(plan.inserts)} · 갱신 {len(plan.updates)} · 변경 없음 {plan.unchanged}"
                      f" · 이번 파일에 없는 기존 행(유지) {plan.db_only}")
                print(f"- 대표 근거: 삽입 {len(plan.ev_insert)} · 갱신 {len(plan.ev_update)} · 변경 없음 {plan.ev_unchanged}"
                      f" · 대표 해제(비대표로 보존) {len(plan.ev_demote)}")
                if plan.issues:
                    print_issues(plan.issues, args.max_errors)
                    trans.rollback()
                    print("결과: 실패 — 재적재 규칙 위반, DB에 쓰지 않음")
                    return 1
                if args.mode == "plan":
                    n, m, digest = fingerprint(conn, vr)
                    print(f"- 이번 파일 항목 중 DB에 있는 행 {n}/{len(vr.records)} · 대표 근거 {m} · PK 지문 {digest}")
                    trans.rollback()
                    print("결과: plan 완료 — DB에 쓰지 않음")
                    return 0
                try:
                    apply_plan(conn, plan)
                except DBAPIError as e:
                    trans.rollback()
                    print(f"DB 오류: {str(e.orig).splitlines()[0]}")
                    print("결과: 실패 — DB가 거부, 사전·근거 전체 롤백")
                    return 1
                post = verify_loaded(conn, vr)
                if post:
                    print_issues(post, args.max_errors)
                    trans.rollback()
                    print("결과: 실패 — 적재 후 검증 불일치, 사전·근거 전체 롤백")
                    return 1
                n, m, digest = fingerprint(conn, vr)
                trans.commit()
                print(f"- 이번 파일 항목 중 DB에 있는 행 {n}/{len(vr.records)} · 대표 근거 {m} · PK 지문 {digest}")
                print(f"결과: 적재 완료 — 사전 삽입 {len(plan.inserts)} · 갱신 {len(plan.updates)}"
                      f" · 대표 근거 삽입 {len(plan.ev_insert)} · 갱신 {len(plan.ev_update)}"
                      f" · 대표 해제 {len(plan.ev_demote)}")
                return 0
            except BaseException:
                if trans.is_active:
                    trans.rollback()
                raise
    finally:
        engine.dispose()


if __name__ == "__main__":
    sys.exit(main())
