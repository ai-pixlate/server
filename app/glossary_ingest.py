"""데이터팀 용어집 3종(xlsx) 검증·적재 CLI.

    python -m app.glossary_ingest validate --certification C.xlsx --ingredient I.xlsx --phrase P.xlsx
    python -m app.glossary_ingest plan     ...   # DB와 비교만 하고 쓰지 않는다
    python -m app.glossary_ingest load     ...   # 전부 검증한 뒤 한 트랜잭션으로 적재

파일은 명시한 경로만 읽는다(하나 이상). 원본 파일은 수정하지 않는다.
DB 접속은 DATABASE_URL 환경변수(plan·load). 접속 정보는 출력하지 않는다.

적재 규칙(데이터팀 9/28 확정 · 사전5종 사용설명서 v3.4.3):
- 대상: 인증·성분은 구역=A, 문구는 구역=A 이고 적재여부=Y. 나머지는 적재하지 않고 사유별로 센다.
  구역은 파일별 안내 시트에 적힌 값만 허용한다(인증 A·B·C · 성분 A·C · 문구 A·C·보류·이관).
  원본 ID 형식·중복은 제외 행까지 모든 데이터 행에서 검사하고, 나머지 값 검사는 적재 대상 행에만 한다.
- 저장 열: 공통 16열 + 문구 3열(적재여부·라벨·대표여부). 열 이름으로 찾고, 그 외 열은 읽지 않는다.
- 값은 원본 그대로 저장한다. 바꾸는 것은 두 가지뿐:
  enforcement 강제→enforced · 참고→reference, 카테고리 'Sunscreens & Tanning'→'Sunscreens & Tanning Products'.
  target_text 의 '; '(대안)·'/'(표현 일부)는 나누지 않는다. 빈 숫자 칸은 0이 아니라 NULL.
- 오류가 한 건이라도 있으면 아무것도 쓰지 않는다(잘못된 행만 건너뛰지 않는다).

재적재(같은 원본 ID = external_id):
- 모든 열이 같으면 변경 없음, version 이 올라갔으면 내부 PK를 유지한 채 갱신.
- 오류로 중단: version 이 낮아짐 · 같은 version 인데 내용이 다름 · 같은 ID인데 자연키
  (target_lang, internal_category, term_ko)가 바뀜(번호 재사용) · 같은 자연키에 다른 원본 ID ·
  원본 ID 없는 기존 행과 자연키가 겹침 · 이미 적재된 ID가 이번 파일에서 적재 제외로 바뀜.
- 이번 파일에 없는 기존 행은 지우지 않고 건수만 보고한다.
- 같은 ID의 한국어·카테고리 정정(오타 수정 포함)도 막는다. 로더로는 반영되지 않으며 별도 검토 절차가 필요하다.

plan 은 읽기 전용 트랜잭션으로 조회만 하고 명시적 쓰기 차단 잠금을 잡지 않는다. load 는 쓰기를 막는 잠금을 잡은 뒤
DB 비교를 처음부터 다시 한다(plan 이후 DB가 바뀌었을 수 있다).
term_ko 길이는 로더가 제한하지 않는다. 자연키 UNIQUE 인덱스에는 항목 크기 한도가 있다(테스트한
PostgreSQL 16 에서 2,704바이트로 관찰. term_ko 단독 한도가 아니며 압축·복합 키 구성에 따라 달라진다).
넘는 값은 load 의 INSERT 에서 DB가 거부하고 전체가 롤백된다 — plan 으로는 미리 알 수 없다.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import os
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from typing import Any, Optional

import openpyxl
from sqlalchemy import create_engine, text
from sqlalchemy.engine import Connection
from sqlalchemy.exc import DBAPIError


@dataclass(frozen=True)
class Spec:
    sheet: str
    id_prefix: str
    term_kind: str
    phrase: bool
    zones: tuple[str, ...]  # 각 파일 00_읽는법에 정의된 구역 값


SPECS = {
    "certification": Spec("용어집_인증", "GL-C", "certification", False, ("A", "B", "C")),
    "ingredient": Spec("용어집_성분", "GL-I", "ingredient", False, ("A", "C")),
    "phrase": Spec("매칭후보", "GL-M", "marketing", True, ("A", "C", "보류", "이관")),
}

COMMON_COLUMNS = [
    "구역", "id", "source_ko", "target_text", "term_kind", "enforcement", "target_lang",
    "internal_category", "example_sentence", "frequency", "corpus_size", "corpus_version",
    "verified_at", "source", "note", "version",
]
PHRASE_COLUMNS = ["적재여부", "라벨", "대표여부"]

ENFORCEMENT = {"강제": "enforced", "참고": "reference"}
CATEGORY_ALIAS = {"Sunscreens & Tanning": "Sunscreens & Tanning Products"}
REPRESENTATIVE = {"대표", "대안"}

# DB 열 → 길이 한도(VARCHAR). TEXT 열은 없음.
VARCHAR_LIMITS = {
    "external_id": 32, "zone": 10, "term_kind": 20, "target_lang": 8, "internal_category": 40,
    "corpus_version": 20, "load_flag": 4, "label": 100, "representative": 10,
}
# glossary 에 쓰는 열(내부 PK 제외). 비교·INSERT·UPDATE 가 이 순서를 쓴다.
DB_FIELDS = [
    "external_id", "term_ko", "term_target", "target_lang", "internal_category", "enforcement",
    "example_sentence", "zone", "term_kind", "frequency", "corpus_size", "corpus_version",
    "verified_at", "source", "note", "version", "load_flag", "label", "representative",
]
NATURAL_KEY = ("target_lang", "internal_category", "term_ko")


@dataclass
class Issue:
    file: str
    row: Optional[int]
    external_id: Optional[str]
    column: Optional[str]
    message: str

    def __str__(self) -> str:
        where = self.file
        if self.row is not None:
            where += f" {self.row}행"
        if self.external_id:
            where += f" {self.external_id}"
        if self.column:
            where += f" [{self.column}]"
        return f"{where}: {self.message}"


@dataclass
class ParsedFile:
    kind: str
    path: str
    total_rows: int = 0
    records: list[dict[str, Any]] = field(default_factory=list)  # 적재 대상, DB 열 이름
    rows: list[int] = field(default_factory=list)  # records 와 같은 순서의 엑셀 행 번호
    excluded: Counter = field(default_factory=Counter)  # 사유 → 건수
    excluded_ids: dict[str, str] = field(default_factory=dict)  # external_id → 사유
    category_normalized: list[str] = field(default_factory=list)
    all_ids: list[tuple[str, int]] = field(default_factory=list)  # 제외 행 포함 모든 행의 (ID, 엑셀 행)


@dataclass
class ValidationResult:
    files: list[ParsedFile]
    issues: list[Issue]

    @property
    def records(self) -> list[dict[str, Any]]:
        return [r for f in self.files for r in f.records]


# ── 셀 값 변환 ────────────────────────────────────────────────────────────

class CellError(ValueError):
    pass


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _text(v: Any, required: bool) -> Optional[str]:
    if _blank(v):
        if required:
            raise CellError("필수값이 비어 있음")
        return None
    if not isinstance(v, str):
        raise CellError(f"문자열이 아님({type(v).__name__}: {v!r})")
    return v  # 원본 그대로(앞뒤 공백 포함) 보존


def _int(v: Any, required: bool) -> Optional[int]:
    if _blank(v):
        if required:
            raise CellError("필수값이 비어 있음")
        return None
    if isinstance(v, bool):
        raise CellError(f"정수가 아님({v!r})")
    if isinstance(v, int):
        n = v
    elif isinstance(v, float) and v.is_integer():
        n = int(v)
    else:
        raise CellError(f"정수가 아님({v!r})")
    if n < 0:
        raise CellError(f"음수({n})")
    return n


def _date(v: Any) -> Optional[dt.date]:
    if _blank(v):
        return None
    if isinstance(v, dt.datetime):
        if v.time() != dt.time(0, 0):
            raise CellError(f"날짜가 아니라 시각이 들어 있음({v.isoformat()})")
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        try:
            return dt.date.fromisoformat(v.strip())
        except ValueError:
            pass
    raise CellError(f"날짜(YYYY-MM-DD)가 아님({v!r})")


# ── 파일 읽기·검증 ────────────────────────────────────────────────────────

def parse_file(kind: str, path: str, issues: list[Issue]) -> ParsedFile:
    spec = SPECS[kind]
    pf = ParsedFile(kind=kind, path=path)
    label = f"{kind}({os.path.basename(path)})"
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # 파일 없음·형식 오류
        issues.append(Issue(label, None, None, None, f"파일을 열 수 없음: {e}"))
        return pf
    try:
        if spec.sheet not in wb.sheetnames:
            issues.append(Issue(label, None, None, None, f"시트 '{spec.sheet}' 없음"))
            return pf
        it = wb[spec.sheet].iter_rows(values_only=True)
        header = [None if h is None else str(h).strip() for h in next(it, ())]
        wanted = COMMON_COLUMNS + (PHRASE_COLUMNS if spec.phrase else [])
        idx: dict[str, int] = {}
        for col in wanted:
            positions = [i for i, h in enumerate(header) if h == col]
            if not positions:
                issues.append(Issue(label, 1, None, col, "필수 열 없음"))
            elif len(positions) > 1:
                issues.append(Issue(label, 1, None, col, "같은 이름의 열이 여러 개"))
            else:
                idx[col] = positions[0]
        if len(idx) != len(wanted):
            return pf

        id_pattern = re.compile(rf"^{re.escape(spec.id_prefix)}\d+$")
        for row_no, row in enumerate(it, start=2):
            if all(_blank(v) for v in row):
                continue
            pf.total_rows += 1
            cell = {c: (row[i] if i < len(row) else None) for c, i in idx.items()}
            _parse_row(pf, spec, label, row_no, cell, id_pattern, issues)
    finally:
        wb.close()
    return pf


def _parse_row(pf, spec, label, row_no, cell, id_pattern, issues) -> None:
    raw_id = cell["id"]
    eid = raw_id if isinstance(raw_id, str) else None
    errs: list[Issue] = []

    def err(col, msg):
        errs.append(Issue(label, row_no, eid, col, msg))

    # 원본 ID는 제외 행까지 모든 행에서 검사한다(번호 재사용·중복 방지).
    if eid is None or not id_pattern.match(eid):
        err("id", f"{spec.id_prefix}+숫자 형식이 아님({raw_id!r})")
    else:
        pf.all_ids.append((eid, row_no))

    zone = cell["구역"]
    if zone not in spec.zones:
        err("구역", f"허용되지 않는 구역 {zone!r} (이 파일은 {'·'.join(spec.zones)})")
        issues.extend(errs)
        return
    load_flag = None
    if spec.phrase:
        load_flag = cell["적재여부"]
        if load_flag not in ("Y", "N"):
            err("적재여부", f"Y/N 이 아님({load_flag!r})")
            issues.extend(errs)
            return
        if load_flag == "Y" and zone != "A":
            err("적재여부", f"구역={zone} 인데 적재여부=Y")
            issues.extend(errs)
            return

    target = zone == "A" and (not spec.phrase or load_flag == "Y")
    if not target:
        # 제외 행은 번역어 등이 비어 있을 수 있다(예: 대응 미정). ID·구역·적재여부만 검사한다.
        issues.extend(errs)
        if errs:
            return
        reason = f"구역={zone}" if zone != "A" else "구역=A·적재여부=N"
        pf.excluded[reason] += 1
        pf.excluded_ids[eid] = reason
        return

    rec: dict[str, Any] = {"zone": zone}

    def take(col, db_col, fn):
        try:
            rec[db_col] = fn(cell[col])
        except CellError as e:
            err(col, str(e))

    rec["external_id"] = eid
    take("source_ko", "term_ko", lambda v: _text(v, True))
    take("target_text", "term_target", lambda v: _text(v, True))
    take("term_kind", "term_kind", lambda v: _text(v, True))
    take("target_lang", "target_lang", lambda v: _text(v, True))
    take("internal_category", "internal_category", lambda v: _text(v, True))
    take("example_sentence", "example_sentence", lambda v: _text(v, False))
    take("frequency", "frequency", lambda v: _int(v, False))
    take("corpus_size", "corpus_size", lambda v: _int(v, False))
    take("corpus_version", "corpus_version", lambda v: _text(v, False))
    take("verified_at", "verified_at", _date)
    take("source", "source", lambda v: _text(v, False))
    take("note", "note", lambda v: _text(v, False))
    take("version", "version", lambda v: _int(v, True))

    enf = cell["enforcement"]
    if enf in ENFORCEMENT:
        rec["enforcement"] = ENFORCEMENT[enf]
    else:
        err("enforcement", f"강제/참고 가 아님({enf!r})")

    if rec.get("term_kind") is not None and rec["term_kind"] != spec.term_kind:
        err("term_kind", f"이 파일은 {spec.term_kind} 인데 {rec['term_kind']!r}")

    cat = rec.get("internal_category")
    if cat in CATEGORY_ALIAS:
        rec["internal_category"] = CATEGORY_ALIAS[cat]
        pf.category_normalized.append(eid or f"{row_no}행")

    if spec.phrase:
        rec["load_flag"] = load_flag
        take("라벨", "label", lambda v: _text(v, True))
        rep = cell["대표여부"]
        if rep in REPRESENTATIVE:
            rec["representative"] = rep
        else:
            err("대표여부", f"대표/대안 이 아님({rep!r})")
    else:
        rec["load_flag"] = rec["label"] = rec["representative"] = None

    for db_col, limit in VARCHAR_LIMITS.items():
        v = rec.get(db_col)
        if isinstance(v, str) and len(v) > limit:
            err(db_col, f"{limit}자 초과({len(v)}자)")

    if errs:
        issues.extend(errs)
        return
    pf.records.append(rec)
    pf.rows.append(row_no)


def validate_files(paths: dict[str, str]) -> ValidationResult:
    issues: list[Issue] = []
    files = [parse_file(kind, path, issues) for kind, path in paths.items()]

    # 파일을 가로질러: 원본 ID 중복(제외 행 포함 전체) · 자연키 중복(적재 대상끼리)
    by_id: dict[str, list[tuple[str, int]]] = defaultdict(list)
    by_key: dict[tuple, list[tuple[str, int, str]]] = defaultdict(list)
    for f in files:
        for eid, row in f.all_ids:
            by_id[eid].append((f.kind, row))
        for rec, row in zip(f.records, f.rows):
            by_key[tuple(rec[k] for k in NATURAL_KEY)].append((f.kind, row, rec["external_id"]))
    for eid, where in by_id.items():
        if len(where) > 1:
            issues.append(Issue("전체", None, eid, "id", f"원본 ID 중복: {where}"))
    for key, where in by_key.items():
        if len(where) > 1:
            issues.append(Issue("전체", None, None, None, f"자연키 중복 {key[:2]}+'{key[2][:30]}': {where}"))
    return ValidationResult(files, issues)


# ── DB 비교·적재 ──────────────────────────────────────────────────────────

@dataclass
class DbPlan:
    inserts: list[dict[str, Any]] = field(default_factory=list)
    updates: list[dict[str, Any]] = field(default_factory=list)  # id(내부 PK) 포함
    unchanged: int = 0
    db_only: int = 0
    issues: list[Issue] = field(default_factory=list)


def _require_schema(conn: Connection) -> None:
    cols = {
        r[0] for r in conn.execute(text(
            "SELECT column_name FROM information_schema.columns "
            "WHERE table_schema = current_schema() AND table_name = 'glossary'"
        ))
    }
    missing = [c for c in DB_FIELDS if c not in cols]
    if missing:
        raise SystemExit(f"glossary 에 열이 없음 {missing} — 마이그레이션 0005 를 먼저 적용")


def build_plan(conn: Connection, vr: ValidationResult) -> DbPlan:
    plan = DbPlan()
    existing = conn.execute(text(f"SELECT id, {', '.join(DB_FIELDS)} FROM glossary")).mappings().all()
    by_ext = {r["external_id"]: r for r in existing if r["external_id"] is not None}
    by_key = {tuple(r[k] for k in NATURAL_KEY): r for r in existing}

    incoming_ids = set()
    for f in vr.files:
        for rec, row in zip(f.records, f.rows):
            eid = rec["external_id"]
            incoming_ids.add(eid)

            def fail(msg):
                plan.issues.append(Issue(f.kind, row, eid, None, msg))

            key = tuple(rec[k] for k in NATURAL_KEY)
            cur = by_ext.get(eid)
            if cur is not None:
                cur_key = tuple(cur[k] for k in NATURAL_KEY)
                if cur_key != key:
                    fail(f"같은 원본 ID인데 자연키가 바뀜(번호 재사용 의심) DB id={cur['id']}")
                    continue
                if all(cur[k] == rec[k] for k in DB_FIELDS):
                    plan.unchanged += 1
                    continue
                cur_v = cur["version"]
                if cur_v is not None and rec["version"] < cur_v:
                    fail(f"version 이 낮아짐(DB {cur_v} → 파일 {rec['version']})")
                elif cur_v is not None and rec["version"] == cur_v:
                    changed = [k for k in DB_FIELDS if cur[k] != rec[k]]
                    fail(f"같은 version({cur_v})인데 내용이 다름 {changed}")
                else:
                    plan.updates.append({**rec, "id": cur["id"]})
                continue
            other = by_key.get(key)
            if other is not None:
                if other["external_id"] is None:
                    fail(f"원본 ID 없는 기존 행(DB id={other['id']})과 자연키가 겹침 — 연결 방법 결정 필요")
                else:
                    fail(f"같은 자연키에 다른 원본 ID가 이미 있음({other['external_id']})")
                continue
            plan.inserts.append(rec)

    # 이미 적재된 ID가 이번 파일에서 적재 제외로 바뀐 경우(강등)
    for f in vr.files:
        for eid, reason in f.excluded_ids.items():
            if eid in by_ext:
                plan.issues.append(Issue(f.kind, None, eid, None,
                                         f"DB에 적재된 행이 이번 파일에서 적재 제외({reason})로 바뀜 — 처리 방식 결정 필요"))
    prefixes = tuple(SPECS[f.kind].id_prefix for f in vr.files)
    plan.db_only = sum(1 for eid in by_ext if eid.startswith(prefixes) and eid not in incoming_ids
                       and not any(eid in f.excluded_ids for f in vr.files))
    return plan


def apply_plan(conn: Connection, plan: DbPlan) -> None:
    cols = ", ".join(DB_FIELDS)
    vals = ", ".join(f":{c}" for c in DB_FIELDS)
    if plan.inserts:
        conn.execute(text(f"INSERT INTO glossary ({cols}) VALUES ({vals})"), plan.inserts)
    if plan.updates:
        sets = ", ".join(f"{c} = :{c}" for c in DB_FIELDS if c != "external_id")
        conn.execute(text(f"UPDATE glossary SET {sets} WHERE id = :id AND external_id = :external_id"), plan.updates)


def verify_loaded(conn: Connection, vr: ValidationResult) -> list[Issue]:
    """적재 직후(커밋 전) 이번 파일의 원본 ID 전부가 파일 값 그대로 들어 있는지 확인한다."""
    issues = []
    rows = conn.execute(
        text(f"SELECT {', '.join(DB_FIELDS)} FROM glossary WHERE external_id = ANY(:ids)"),
        {"ids": [r["external_id"] for r in vr.records]},
    ).mappings().all()
    db = {r["external_id"]: r for r in rows}
    for rec in vr.records:
        cur = db.get(rec["external_id"])
        if cur is None:
            issues.append(Issue("검증", None, rec["external_id"], None, "적재 후 행이 없음"))
        elif any(cur[k] != rec[k] for k in DB_FIELDS):
            issues.append(Issue("검증", None, rec["external_id"], None, "적재 후 값이 파일과 다름"))
    return issues


def file_fingerprint(conn: Connection, vr: ValidationResult) -> tuple[int, str]:
    """이번 파일의 원본 ID 중 DB에 있는 행 수와 (원본 ID:내부 PK) 목록의 md5.

    적재 직후와 나중 plan 의 값이 같으면 이번 파일 행들의 PK가 그대로라는 뜻이다.
    다른 전달본의 행은 섞이지 않는다."""
    rows = conn.execute(
        text("SELECT external_id, id FROM glossary WHERE external_id = ANY(:ids)"),
        {"ids": [r["external_id"] for r in vr.records]},
    ).all()
    joined = ",".join(f"{e}:{i}" for e, i in sorted(rows))
    return len(rows), hashlib.md5(joined.encode("utf-8")).hexdigest()


# ── 출력 ─────────────────────────────────────────────────────────────────

def print_summary(vr: ValidationResult, out=None) -> None:
    out = out or sys.stdout
    print("== 파일 검증 ==", file=out)
    for f in vr.files:
        excl = ", ".join(f"{k} {v}" for k, v in sorted(f.excluded.items())) or "없음"
        print(f"- {f.kind}: 데이터 {f.total_rows}행 · 적재 대상 {len(f.records)}행 · 제외 {sum(f.excluded.values())}행 ({excl})", file=out)
        a_n = sorted(e for e, r in f.excluded_ids.items() if r == "구역=A·적재여부=N")
        if a_n:
            print(f"    구역=A·적재여부=N: {', '.join(a_n)}", file=out)
        if f.category_normalized:
            print(f"    카테고리 정규화 {len(f.category_normalized)}행: {', '.join(f.category_normalized)}", file=out)
    recs = vr.records
    long_rows = sum(1 for r in recs if len(r["term_ko"]) > 200 or len(r["term_target"]) > 200)
    cats = Counter(r["internal_category"] for r in recs)
    print(f"- 합계 적재 대상 {len(recs)}행 · 200자 초과(원본 보존) {long_rows}행 · 카테고리 {dict(cats)}", file=out)


def print_issues(issues: list[Issue], limit: int, out=None) -> None:
    out = out or sys.stdout
    print(f"== 오류 {len(issues)}건 ==", file=out)
    for i in issues[:limit]:
        print(f"  {i}", file=out)
    if len(issues) > limit:
        print(f"  … 외 {len(issues) - limit}건 (--max-errors 로 더 보기)", file=out)


# ── CLI ──────────────────────────────────────────────────────────────────

def main(argv: Optional[list[str]] = None, database_url: Optional[str] = None) -> int:
    p = argparse.ArgumentParser(prog="python -m app.glossary_ingest", description="용어집 xlsx 검증·적재")
    p.add_argument("mode", choices=["validate", "plan", "load"],
                   help="validate=파일만 검사(DB 없음) · plan=DB와 비교만(쓰기 없음) · load=적재")
    for kind in SPECS:
        p.add_argument(f"--{kind}", metavar="XLSX", help=f"{SPECS[kind].sheet} 시트가 있는 파일")
    p.add_argument("--max-errors", type=int, default=50, help="출력할 오류 수(기본 50)")
    args = p.parse_args(argv)

    paths = {k: getattr(args, k) for k in SPECS if getattr(args, k)}
    if not paths:
        p.error("--certification / --ingredient / --phrase 중 하나 이상 필요")

    vr = validate_files(paths)
    print_summary(vr)
    if vr.issues:
        print_issues(vr.issues, args.max_errors)
        print("결과: 실패 — 파일 오류가 있어 DB에 쓰지 않음")
        return 1
    if args.mode == "validate":
        print("결과: 파일 검증 통과 (DB 미접속)")
        return 0

    url = database_url or os.getenv("DATABASE_URL")
    if not url:
        print("DATABASE_URL 환경변수가 없음")
        return 2
    engine = create_engine(url, future=True)
    try:
        with engine.connect() as conn:
            trans = conn.begin()
            try:
                if args.mode == "plan":
                    # 조회만 한다. 명시적 쓰기 차단 잠금을 잡지 않는다(일반 조회의 읽기 잠금만).
                    conn.execute(text("SET TRANSACTION READ ONLY"))
                _require_schema(conn)
                if args.mode == "load":
                    # 다른 쓰기와 동시에 돌지 않게 쓰기만 막는다(읽기는 허용). 비교는 잠금 후 처음부터 다시 한다.
                    conn.execute(text("SET LOCAL lock_timeout = '10s'"))
                    conn.execute(text("LOCK TABLE glossary IN SHARE ROW EXCLUSIVE MODE"))
                plan = build_plan(conn, vr)
                print("== DB 비교 ==")
                print(f"- 삽입 {len(plan.inserts)} · 갱신 {len(plan.updates)} · 변경 없음 {plan.unchanged}"
                      f" · 이번 파일에 없는 기존 행(유지) {plan.db_only}")
                if plan.issues:
                    print_issues(plan.issues, args.max_errors)
                    trans.rollback()
                    print("결과: 실패 — 재적재 규칙 위반, DB에 쓰지 않음")
                    return 1
                if args.mode == "plan":
                    present, digest = file_fingerprint(conn, vr)
                    print(f"- 이번 파일 원본 ID 중 DB에 있는 행 {present}/{len(vr.records)} · PK 지문 {digest}")
                    trans.rollback()
                    print("결과: plan 완료 — DB에 쓰지 않음")
                    return 0
                try:
                    apply_plan(conn, plan)
                except DBAPIError as e:
                    trans.rollback()
                    print(f"DB 오류: {str(e.orig).splitlines()[0]}")
                    print("결과: 실패 — DB가 거부, 전체 롤백")
                    return 1
                post = verify_loaded(conn, vr)
                if post:
                    print_issues(post, args.max_errors)
                    trans.rollback()
                    print("결과: 실패 — 적재 후 검증 불일치, 롤백")
                    return 1
                present, digest = file_fingerprint(conn, vr)
                trans.commit()
                print(f"- 이번 파일 원본 ID 중 DB에 있는 행 {present}/{len(vr.records)} · PK 지문 {digest}")
                print(f"결과: 적재 완료 — 삽입 {len(plan.inserts)} · 갱신 {len(plan.updates)} · 변경 없음 {plan.unchanged}")
                return 0
            except BaseException:
                if trans.is_active:
                    trans.rollback()
                raise
    finally:
        engine.dispose()


if __name__ == "__main__":
    sys.exit(main())
