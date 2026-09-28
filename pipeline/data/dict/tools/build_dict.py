"""전달 자료 xlsx → 정규화 사전 JSON 변환 · 검증 · 비교 도구.

근거: docs/ai-experiments/2026-09-28_03-1-judge_design-v1.md 8절(D8), pipeline/dictionary.py(스키마).
실제 사전 값은 git 밖(`judge.dict_dir`, 기본 pipeline/samples/local/dict/normalized)에 쓴다. 이 도구와 스키마만 git에 있다.

    python -m pipeline.data.dict.tools.build_dict build \
        --regulation <regulation_dict.xlsx> --local <locale_unsuitable_dict.xlsx> --out <dir> \
        [--date 2026-09-28] [--seq 1] [--local-version-status mismatch_pending --local-version-note "..."]
    python -m pipeline.data.dict.tools.build_dict verify --dir <dir>
    python -m pipeline.data.dict.tools.build_dict diff --old <dir> --new <dir>

허용 열(색이 아니라 목록으로 정한다, D8):
- 규제사전 탭: REG_COLUMNS. `confidence`는 검증(high만 적재)에만 쓰고 출력하지 않는다. `source` · `version`은 시트 전용이라 읽지 않는다.
- 현지부적합 탭: LOCAL_COLUMNS. `판단 근거` · `kr_freq` · `kr_corpus`는 내부 기록용(설명서 §3)이라 읽지 않는다.
- 근거 탭: 대표 근거(`is_primary=Y`, `rg_id`에 해당 id 포함)의 quote · verified_at을 가져오고 article · url이 규제사전 행과 같은지 검사한다.

policy_rules.json은 xlsx에 없다. 설명서 §3 · §4의 분기표를 이 파일의 RULES_TEMPLATE(잠정, D10)에서 만들며 overrides는 빈 목록이다(D2 — 실제 예외 쌍은
데이터 담당 확인 후에만 채운다. 이 도구는 채우지 않는다).
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path
from typing import Any

from pipeline.dictionary import (
    DICT_FILES,
    DICT_SCHEMA_VERSION,
    LOCAL_FILE,
    REGULATION_FILE,
    RULES_FILE,
    DictionaryError,
    LocalDict,
    PolicyRules,
    RegulationDict,
    dump_model,
    load_dictionaries,
    sha256_of,
)

REG_SHEET = "규제사전"
EVIDENCE_SHEET = "근거"
LOCAL_SHEET = "01_현지부적합사전"
README_SHEET = "00_읽는법"
SUMS_FILE = "SHA256SUMS"

REG_COLUMNS = (
    "id", "dict_type", "target_country", "regulatory_class", "source_expression",
    "variant_expressions.ko", "variant_expressions.en", "alternative_expression", "verdict_status", "reason",
    "evidence_id", "evidence_source_type", "evidence_article", "evidence_url", "verified_at", "internal_category",
)
REG_VALIDATION_ONLY = ("confidence",)
LOCAL_COLUMNS = ("id", "항목", "패턴", "판정", "제외하는 맥락", "제외하지 않는 맥락", "셀러 문장", "verified_at")
EVIDENCE_COLUMNS = ("evidence_id", "rg_id", "source_type", "company", "issued_at", "quote", "article", "url", "is_primary", "verified_at")
SEP = "; "

RULES_TEMPLATE: dict[str, Any] = {
    "schema_version": DICT_SCHEMA_VERSION,
    "rules_version": None,  # build 시 policy@<date>.<seq>
    "basis": (
        "사전 5종 사용 설명서 v3.4.3 §3 · §4 기반 잠정(open-questions #7 · #60, 설계 v1 D10). "
        "저장소 계약에 채택된 것이 아니며 원문(결정기록 · PRD F-SEC-04) 확인 대기. overrides는 근거 확인 전 빈 목록(D2)."
    ),
    "regulatory_class_map": {
        "cosmetic": {"applied_classes": ["cosmetic", "common"], "dict_coverage": "selected_class_all_entries"},
        "otc": {"applied_classes": ["otc", "common"], "dict_coverage": "selected_class_all_entries"},
        "combination": {"applied_classes": ["otc", "common"], "dict_coverage": "partial_class_combination"},
        "unknown": {"applied_classes": ["cosmetic", "common"], "dict_coverage": "unverified_class"},
    },
    "verdict_map": [
        {"dict_type": "regulatory", "verdict_status": "allowed", "alternative": "any", "emit_verdict_status": None, "bucket": "none"},
        {"dict_type": "regulatory", "verdict_status": "conditional", "alternative": "any", "emit_verdict_status": "conditional", "bucket": "include"},
        {"dict_type": "regulatory", "verdict_status": "rewritable", "alternative": "required", "emit_verdict_status": "regulated", "bucket": "include"},
        {"dict_type": "regulatory", "verdict_status": "regulated", "alternative": "present", "emit_verdict_status": "regulated", "bucket": "include"},
        {"dict_type": "regulatory", "verdict_status": "regulated", "alternative": "absent", "emit_verdict_status": "regulated", "bucket": "exclude"},
        {"dict_type": "local", "verdict_status": "irrelevant", "alternative": "any", "emit_verdict_status": "irrelevant", "bucket": "exclude"},
        {"dict_type": "local", "verdict_status": "needs_fix", "alternative": "any", "emit_verdict_status": "needs_fix", "bucket": "exclude"},
    ],
    "uncertain_bucket": "exclude",  # D12 추천(잠정) — 제외 권고 + 확신 상태는 별도 필드
    "overrides": [],
}


class BuildError(ValueError):
    pass


# ---------------------------------------------------------------------------
# xlsx 읽기
# ---------------------------------------------------------------------------
def _load_workbook(path: Path):
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise BuildError("openpyxl이 필요하다: pip install openpyxl") from e
    return openpyxl.load_workbook(path, data_only=True, read_only=True)


def _sheet_rows(wb, name: str) -> list[list[Any]]:
    if name not in wb.sheetnames:
        raise BuildError(f"탭이 없다: {name!r} (있는 탭: {wb.sheetnames})")
    rows = []
    for r in wb[name].iter_rows(values_only=True):
        if any(c is not None and str(c).strip() for c in r):
            rows.append(list(r))
    return rows


def _table(rows: list[list[Any]], sheet: str) -> list[dict[str, Any]]:
    """1행 = 열 이름(규제사전 00_읽는법 규칙). 빈 열 이름은 무시한다."""
    if not rows:
        raise BuildError(f"{sheet}: 비어 있다")
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = []
    for r in rows[1:]:
        rec: dict[str, Any] = {}
        for k, v in zip(header, r):
            if k:
                rec[k] = _cell(v)
        out.append(rec)
    return out


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (dt.date, dt.datetime)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _multi(v: str) -> list[str]:
    return [p.strip() for p in v.split(SEP.strip()) if p.strip()] if v else []


def _require(rec: dict[str, Any], cols: tuple[str, ...], sheet: str) -> None:
    missing = [c for c in cols if c not in rec]
    if missing:
        raise BuildError(f"{sheet}: 없는 열 {missing}")


def _claimed_version(readme_rows: list[list[Any]]) -> str | None:
    """00_읽는법에서 원본이 밝힌 버전. 규제사전은 'YYYY-MM-DD vN' 이력 중 최신 N, 현지부적합은 'vN · 대조일 …' 문구."""
    best: tuple[int, str] | None = None
    for r in readme_rows:
        for c in r:
            if c is None:
                continue
            s = str(c)
            m = re.match(r"^(\d{4}-\d{2}-\d{2}) v(\d+)$", s.strip())
            if m:
                n = int(m.group(2))
                if best is None or n > best[0]:
                    best = (n, f"v{n} ({m.group(1)})")
            m2 = re.match(r"^v(\d+) · 대조일 (\d{4}-\d{2}-\d{2})", s.strip())
            if m2:
                return f"v{m2.group(1)} (대조일 {m2.group(2)})"
    return best[1] if best else None


# ---------------------------------------------------------------------------
# 규제사전
# ---------------------------------------------------------------------------
def build_regulation(xlsx: Path, *, date: str, seq: int) -> RegulationDict:
    wb = _load_workbook(xlsx)
    reg = _table(_sheet_rows(wb, REG_SHEET), REG_SHEET)
    ev = _table(_sheet_rows(wb, EVIDENCE_SHEET), EVIDENCE_SHEET)
    claimed = _claimed_version(_sheet_rows(wb, README_SHEET))
    if reg:
        _require(reg[0], REG_COLUMNS + REG_VALIDATION_ONLY, REG_SHEET)
    if ev:
        _require(ev[0], EVIDENCE_COLUMNS, EVIDENCE_SHEET)
    primary: dict[str, list[dict[str, Any]]] = {}
    for e in ev:
        if e["is_primary"].upper() != "Y":
            continue
        for rid in _multi(e["rg_id"]):
            primary.setdefault(rid, []).append(e)

    errors: list[str] = []
    entries: list[dict[str, Any]] = []
    for rec in reg:
        rid = rec["id"]
        if rec["confidence"] != "high":
            errors.append(f"{rid}: confidence={rec['confidence']!r} — 9월 시드는 high만 적재(설명서 §4). 보류 탭에 있어야 한다")
            continue
        for col in ("verified_at", "evidence_id", "evidence_url", "reason"):
            if not rec[col]:
                errors.append(f"{rid}: {col}이 비어 있다")
        cands = [e for e in primary.get(rid, []) if e["evidence_id"] == rec["evidence_id"]]
        if not cands:
            errors.append(f"{rid}: 근거 탭에 evidence_id={rec['evidence_id']!r} · is_primary=Y · rg_id 포함 행이 없다")
            quote = None
            ev_verified = None
        else:
            e = cands[0]
            if e["article"] != rec["evidence_article"] or e["url"] != rec["evidence_url"]:
                errors.append(f"{rid}: 규제사전의 evidence_article/url이 근거 탭과 다르다")
            quote = e["quote"] or None
            ev_verified = e["verified_at"] or None
        entries.append(
            {
                "id": rid,
                "dict_type": rec["dict_type"],
                "target_country": rec["target_country"],
                "regulatory_class": rec["regulatory_class"],
                "source_expression": rec["source_expression"],
                "variant_expressions.ko": _multi(rec["variant_expressions.ko"]),
                "variant_expressions.en": _multi(rec["variant_expressions.en"]),
                "alternative_expression": _multi(rec["alternative_expression"]),
                "verdict_status": rec["verdict_status"],
                "reason": rec["reason"],
                "evidence": {
                    "evidence_id": rec["evidence_id"],
                    "source_type": rec["evidence_source_type"],
                    "article": rec["evidence_article"] or None,
                    "url": rec["evidence_url"],
                    "quote": quote,
                    "verified_at": ev_verified,
                },
                "verified_at": rec["verified_at"],
                "internal_category": rec["internal_category"] or None,
            }
        )
    if errors:
        raise BuildError(f"{REG_SHEET}: " + "; ".join(errors))
    payload = {
        "schema_version": DICT_SCHEMA_VERSION,
        "dictionary_version": f"regulation@{date}.{seq}",
        "dict_type": "regulatory",
        "source": {
            "file": xlsx.name,
            "sha256": sha256_of(xlsx),
            "sheets": [REG_SHEET, EVIDENCE_SHEET],
            "claimed_version": claimed,
            "version_status": "as_claimed" if claimed else "unknown",
            "version_note": None,
            "extracted_at": date,
            "row_count": len(entries),
        },
        "entries": entries,
    }
    try:
        return RegulationDict.model_validate(payload)
    except ValueError as e:
        raise BuildError(f"{REG_SHEET}: {e}") from e


# ---------------------------------------------------------------------------
# 현지부적합사전
# ---------------------------------------------------------------------------
def build_local(xlsx: Path, *, date: str, seq: int, version_status: str, version_note: str | None) -> LocalDict:
    wb = _load_workbook(xlsx)
    rows = _table(_sheet_rows(wb, LOCAL_SHEET), LOCAL_SHEET)
    claimed = _claimed_version(_sheet_rows(wb, README_SHEET))
    if rows:
        _require(rows[0], LOCAL_COLUMNS, LOCAL_SHEET)
    errors: list[str] = []
    entries = []
    for rec in rows:
        rid = rec["id"]
        for col in LOCAL_COLUMNS:
            if not rec[col]:
                errors.append(f"{rid}: {col}이 비어 있다")
        entries.append(
            {
                "id": rid,
                "항목": rec["항목"],
                "패턴": _multi(rec["패턴"]),
                "판정": rec["판정"],
                "제외하는 맥락": rec["제외하는 맥락"],
                "제외하지 않는 맥락": rec["제외하지 않는 맥락"],
                "셀러 문장": rec["셀러 문장"],
                "verified_at": rec["verified_at"],
            }
        )
    if errors:
        raise BuildError(f"{LOCAL_SHEET}: " + "; ".join(errors))
    payload = {
        "schema_version": DICT_SCHEMA_VERSION,
        "dictionary_version": f"local@{date}.{seq}",
        "dict_type": "local",
        "target_country": "US",
        "source": {
            "file": xlsx.name,
            "sha256": sha256_of(xlsx),
            "sheets": [LOCAL_SHEET],
            "claimed_version": claimed,
            "version_status": version_status,
            "version_note": version_note,
            "extracted_at": date,
            "row_count": len(entries),
        },
        "entries": entries,
    }
    try:
        return LocalDict.model_validate(payload)
    except ValueError as e:
        raise BuildError(f"{LOCAL_SHEET}: {e}") from e


def build_rules(*, date: str, seq: int) -> PolicyRules:
    raw = json.loads(json.dumps(RULES_TEMPLATE))
    raw["rules_version"] = f"policy@{date}.{seq}"
    return PolicyRules.model_validate(raw)


# ---------------------------------------------------------------------------
# 묶음 쓰기 · 검증 · 비교
# ---------------------------------------------------------------------------
def write_bundle(out: Path, reg: RegulationDict, loc: LocalDict, rules: PolicyRules) -> dict[str, str]:
    out.mkdir(parents=True, exist_ok=True)
    (out / REGULATION_FILE).write_text(dump_model(reg), encoding="utf-8")
    (out / LOCAL_FILE).write_text(dump_model(loc), encoding="utf-8")
    (out / RULES_FILE).write_text(dump_model(rules), encoding="utf-8")
    sums = {name: sha256_of(out / name) for name in DICT_FILES}
    (out / SUMS_FILE).write_text("".join(f"{h}  {n}\n" for n, h in sums.items()), encoding="utf-8")
    return sums


def verify_bundle(d: Path) -> dict[str, Any]:
    """SHA256SUMS 대조 + 스키마 검증(load_dictionaries). 어긋나면 DictionaryError."""
    sums_path = d / SUMS_FILE
    if not sums_path.exists():
        raise DictionaryError(f"{SUMS_FILE}가 없다: {d}")
    expected: dict[str, str] = {}
    for line in sums_path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            h, n = line.split(None, 1)
            expected[n.strip()] = h
    bad = [n for n in DICT_FILES if expected.get(n) != sha256_of(d / n)]
    if bad:
        raise DictionaryError(f"해시 불일치: {bad}")
    dicts = load_dictionaries(d)
    return {"dir": str(d), "dictionary_version": dicts.dictionary_version, "fingerprint": dicts.fingerprint}


def diff_bundles(old: Path, new: Path) -> dict[str, Any]:
    """어떤 종류의 열이 바뀌었는지 요약(설계 1절 재실행 조건표 · #49). 패턴 / 맥락 / 정책 필드 / 규칙."""
    a, b = load_dictionaries(old), load_dictionaries(new)
    summary: dict[str, list[str]] = {"pattern": [], "context": [], "policy_field": [], "rules": [], "added": [], "removed": []}
    ra = {e.id: e for e in a.regulation.entries} | {e.id: e for e in a.local.entries}
    rb = {e.id: e for e in b.regulation.entries} | {e.id: e for e in b.local.entries}
    summary["added"] = sorted(set(rb) - set(ra))
    summary["removed"] = sorted(set(ra) - set(rb))
    for rid in sorted(set(ra) & set(rb)):
        x, y = ra[rid].model_dump(by_alias=True), rb[rid].model_dump(by_alias=True)
        for k in sorted(set(x) | set(y)):
            if x.get(k) == y.get(k):
                continue
            if k in ("variant_expressions.ko", "variant_expressions.en", "패턴"):
                summary["pattern"].append(f"{rid}.{k}")
            elif k in ("제외하는 맥락", "제외하지 않는 맥락"):
                summary["context"].append(f"{rid}.{k}")
            else:
                summary["policy_field"].append(f"{rid}.{k}")
    if a.rules.model_dump() != b.rules.model_dump():
        summary["rules"].append(f"{a.rules.rules_version} → {b.rules.rules_version}")
    return summary


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build")
    b.add_argument("--regulation", required=True, type=Path)
    b.add_argument("--local", required=True, type=Path)
    b.add_argument("--out", required=True, type=Path)
    b.add_argument("--date", default=dt.date.today().isoformat())
    b.add_argument("--seq", type=int, default=1)
    b.add_argument("--local-version-status", default="as_claimed", choices=["as_claimed", "mismatch_pending", "unknown"])
    b.add_argument("--local-version-note", default=None)
    v = sub.add_parser("verify")
    v.add_argument("--dir", required=True, type=Path)
    d = sub.add_parser("diff")
    d.add_argument("--old", required=True, type=Path)
    d.add_argument("--new", required=True, type=Path)
    a = p.parse_args(argv)
    try:
        if a.cmd == "build":
            reg = build_regulation(a.regulation, date=a.date, seq=a.seq)
            loc = build_local(a.local, date=a.date, seq=a.seq, version_status=a.local_version_status, version_note=a.local_version_note)
            rules = build_rules(date=a.date, seq=a.seq)
            sums = write_bundle(a.out, reg, loc, rules)
            print(json.dumps({"out": str(a.out), "regulation": len(reg.entries), "local": len(loc.entries), "sha256": sums,
                              "dictionary_version": {"regulation": reg.dictionary_version, "local": loc.dictionary_version, "rules": rules.rules_version}},
                             ensure_ascii=False, indent=2))
        elif a.cmd == "verify":
            print(json.dumps(verify_bundle(a.dir), ensure_ascii=False, indent=2))
        else:
            print(json.dumps(diff_bundles(a.old, a.new), ensure_ascii=False, indent=2))
    except (BuildError, DictionaryError) as e:
        print(f"오류: {e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
