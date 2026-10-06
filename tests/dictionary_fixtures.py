"""규제사전·현지부적합 테스트용 xlsx 생성 도우미. 실제 파일과 같은 탭·열 이름으로 작은 파일을 만든다."""
from pathlib import Path

import openpyxl

from app.dictionary_ingest import EVIDENCE_COLUMNS, EVIDENCE_SHEET, LOCAL_SHEET, REG_COLUMNS, REG_SHEET

URL_WL = "https://www.fda.gov/warning-letters/example"
URL_MN = "https://www.accessdata.fda.gov/monograph/M020"
URL_MP = "https://www.amazon.com/dp/EXAMPLE"


def reg_row(n: int, evidence: str = "WL-001", **over) -> dict:
    kind = evidence[:2]
    row = {
        "id": f"RG-{n:03d}",
        "dict_type": "regulatory",
        "target_country": "US",
        "regulatory_class": "cosmetic",
        "source_expression": f"표현{n}",
        "variant_expressions.ko": f"표현{n}; 변형{n}",
        "variant_expressions.en": f"claim {n}; claims {n}",
        "alternative_expression": f"safe {n}",
        "verdict_status": "regulated",
        "reason": "사유 문장",
        "evidence_id": evidence,
        "evidence_source_type": {"WL": "Warning Letter", "MN": "OTC 모노그래프", "MP": "시장 관행"}[kind],
        "evidence_article": None if kind == "MP" else "201(g)(1)",
        "evidence_url": {"WL": URL_WL, "MN": URL_MN, "MP": URL_MP}[kind],
        "verified_at": "2026-09-16",
        "internal_category": None,
        "confidence": "high",
        "source": "테스트",
        "version": "1",
    }
    row.update(over)
    return row


def ev_row(evidence: str, rg_ids: str, primary: str = "Y", **over) -> dict:
    kind = evidence[:2]
    row = {
        "evidence_id": evidence,
        "rg_id": rg_ids,
        "source_type": {"WL": "Warning Letter", "MN": "OTC 모노그래프", "MP": "시장 관행", "GD": "FDA 공식 문서"}[kind],
        "company": None,
        "issued_at": None,
        "quote": f"quote of {evidence}",
        "article": None if kind == "MP" else "201(g)(1)",
        "url": {"WL": URL_WL, "MN": URL_MN, "MP": URL_MP, "GD": URL_WL}[kind],
        "is_primary": primary,
        "verified_at": "2026-09-15",
    }
    row.update(over)
    return row


def local_row(n: int, **over) -> dict:
    row = {
        "id": f"LC-{n:02d}",
        "항목": f"항목{n}",
        "패턴": f"패턴{n}; 패턴{n}b",
        "판정": "irrelevant",
        "제외하는 맥락": f"제외 맥락 {n}",
        "제외하지 않는 맥락": f"유지 맥락 {n}",
        "셀러 문장": "셀러에게 보이는 문장",
        "판단 근거": "내부 기록",
        "kr_freq": 10,
        "kr_corpus": 66,
        "verified_at": "2026-09-22",
    }
    row.update(over)
    return row


def _sheet(wb, name, header, rows):
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append([r.get(h) for h in header])


def write_regulatory(path: Path, rows: list[dict], evidence: list[dict]) -> Path:
    wb = openpyxl.Workbook()
    wb.active.title = "00_읽는법"
    _sheet(wb, REG_SHEET, REG_COLUMNS, rows)
    _sheet(wb, EVIDENCE_SHEET, EVIDENCE_COLUMNS, evidence)
    _sheet(wb, "보류", REG_COLUMNS + ["hold_reason"], [])  # 로더가 읽지 않는 탭
    wb.save(path)
    return path


def write_local(path: Path, rows: list[dict]) -> Path:
    wb = openpyxl.Workbook()
    wb.active.title = "00_읽는법"
    header = ["id", "항목", "패턴", "판정", "제외하는 맥락", "제외하지 않는 맥락", "셀러 문장", "판단 근거",
              "kr_freq", "kr_corpus", "verified_at"]
    _sheet(wb, LOCAL_SHEET, header, rows)
    wb.save(path)
    return path


def basic_regulatory(path: Path) -> Path:
    """RG-001 교체형(WL-001) · RG-002 완충형(WL-002) · RG-003/004 가 MN-001 공유 · RG-005 allowed(MP-001)."""
    rows = [
        reg_row(1, "WL-001"),
        reg_row(2, "WL-002", verdict_status="rewritable"),
        reg_row(3, "MN-001", regulatory_class="otc", alternative_expression=None),
        reg_row(4, "MN-001", regulatory_class="otc", verdict_status="conditional", alternative_expression="SPF 30"),
        reg_row(5, "MP-001", verdict_status="allowed", alternative_expression=None),
    ]
    evidence = [
        ev_row("WL-001", "RG-001"),
        ev_row("WL-002", "RG-002"),
        ev_row("WL-009", "RG-001; RG-002", primary="N"),
        ev_row("MN-001", "RG-003; RG-004"),
        ev_row("MP-001", "RG-005"),
        ev_row("WL-050", "RG-099"),  # 탭에 없는 항목(보류)만 가리킴 — 읽지 않음
    ]
    return write_regulatory(path, rows, evidence)
