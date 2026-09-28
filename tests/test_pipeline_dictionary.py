"""③-1 · ③-1' 사전 데이터 — 스키마 · 로더 · 두 뷰 · 변환 도구(합성 데이터와 임시 xlsx만 쓴다. 실제 사전 값은 쓰지 않는다)."""
import json
from pathlib import Path

import pytest
from pydantic import ValidationError

from pipeline import dictionary as D
from pipeline.data.dict.tools import build_dict as B

SYNTH = Path(__file__).resolve().parent.parent / "pipeline" / "data" / "dict" / "synthetic"


# ---------------------------------------------------------------------------
# 로더 · 두 뷰
# ---------------------------------------------------------------------------
def test_synthetic_bundle_loads_and_verifies():
    dicts = D.load_dictionaries(SYNTH)
    assert len(dicts.regulation.entries) == 6
    assert len(dicts.local.entries) == 3
    assert dicts.dictionary_version == {
        "regulation": "regulation@2026-09-28.0",
        "local": "local@2026-09-28.0",
        "rules": "policy@2026-09-28.0",
    }
    assert set(dicts.fingerprint) == set(D.DICT_FILES)
    assert all(len(h) == 64 for h in dicts.fingerprint.values())
    out = B.verify_bundle(SYNTH)
    assert out["fingerprint"] == dicts.fingerprint


def test_judge_view_exposes_no_policy_fields():
    view = D.load_dictionaries(SYNTH).judge_view()
    assert {i.id for i in view.local_items()} == {"LC-91", "LC-92", "LC-93"}
    assert len(view.regulatory_items()) == 6  # allowed 행(RG-905)도 ③-1은 그대로 본다(정책값을 모른다)
    item = view.by_id("LC-93")
    assert item.patterns_ko == ("원", "가짜만원")
    assert item.exclusion_context and item.keep_context
    rg = view.by_id("RG-902")
    assert rg.exclusion_context is None and rg.keep_context is None
    for attr in ("verdict_status", "alternative_expression", "reason", "regulatory_class", "seller_message", "evidence"):
        assert not hasattr(item, attr) and not hasattr(rg, attr)
    assert view.dictionary_version["local"] == "local@2026-09-28.0"


def test_policy_view_exposes_rules_and_entries():
    pv = D.load_dictionaries(SYNTH).policy_view()
    assert pv.regulation["RG-902"].verdict_status == "regulated" and pv.regulation["RG-902"].alternative_expression == []
    assert pv.regulation["RG-904"].verdict_status == "rewritable" and pv.regulation["RG-904"].alternative_expression
    assert pv.local["LC-92"].verdict_status == "needs_fix" and pv.local["LC-92"].seller_message
    assert set(pv.rules.regulatory_class_map) == {"cosmetic", "otc", "combination", "unknown"}
    assert pv.rules.regulatory_class_map["combination"].dict_coverage == "partial_class_combination"
    assert pv.rules.regulatory_class_map["unknown"].dict_coverage == "unverified_class"
    assert pv.rules.uncertain_bucket == "exclude"
    assert [o.keep for o in pv.rules.overrides] == ["RG-903"]  # 합성 예외 쌍(형식 검증용)


def test_loader_rejects_missing_file_and_wrong_schema_version(tmp_path):
    with pytest.raises(D.DictionaryError, match="사전 파일이 없다"):
        D.load_dictionaries(tmp_path)
    for name in D.DICT_FILES:
        (tmp_path / name).write_bytes((SYNTH / name).read_bytes())  # 바이트 복사 — 줄 끝(LF)을 보존해야 해시가 맞는다
    raw = json.loads((tmp_path / D.LOCAL_FILE).read_text(encoding="utf-8"))
    raw["schema_version"] = "9"
    (tmp_path / D.LOCAL_FILE).write_text(json.dumps(raw, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(D.DictionaryError, match="schema_version"):
        D.load_dictionaries(tmp_path)


def test_loader_rejects_override_pointing_to_unknown_entry(tmp_path):
    for name in D.DICT_FILES:
        (tmp_path / name).write_bytes((SYNTH / name).read_bytes())  # 바이트 복사 — 줄 끝(LF)을 보존해야 해시가 맞는다
    raw = json.loads((tmp_path / D.RULES_FILE).read_text(encoding="utf-8"))
    raw["overrides"] = [{"keep": "RG-903", "drop": ["RG-999"], "basis": "x"}]
    (tmp_path / D.RULES_FILE).write_text(json.dumps(raw, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(D.DictionaryError, match="RG-999"):
        D.load_dictionaries(tmp_path)


def test_verify_detects_tampering(tmp_path):
    for name in (*D.DICT_FILES, B.SUMS_FILE):
        (tmp_path / name).write_bytes((SYNTH / name).read_bytes())  # 바이트 복사 — 줄 끝(LF)을 보존해야 해시가 맞는다
    B.verify_bundle(tmp_path)
    p = tmp_path / D.LOCAL_FILE
    p.write_text(p.read_text(encoding="utf-8").replace("합성 안내문 1", "바뀐 안내문"), encoding="utf-8")
    with pytest.raises(D.DictionaryError, match="해시 불일치"):
        B.verify_bundle(tmp_path)


# ---------------------------------------------------------------------------
# 스키마 규칙(설명서 §4 · 원본 열 이름 별칭)
# ---------------------------------------------------------------------------
def _reg_entry(**over):
    base = {
        "id": "RG-910", "dict_type": "regulatory", "target_country": "US", "regulatory_class": "cosmetic",
        "source_expression": "x", "variant_expressions.ko": ["x"], "variant_expressions.en": [],
        "alternative_expression": [], "verdict_status": "regulated", "reason": "r",
        "evidence": {"evidence_id": "WL-1", "source_type": "Warning Letter", "article": None, "url": "https://e", "quote": None, "verified_at": None},
        "verified_at": "2026-09-28", "internal_category": None,
    }
    base.update(over)
    return base


def test_regulation_entry_alias_and_alternative_rules():
    e = D.RegulationEntry.model_validate(_reg_entry())
    assert e.variant_ko == ["x"] and e.model_dump(by_alias=True)["variant_expressions.ko"] == ["x"]
    with pytest.raises(ValidationError, match="rewritable"):
        D.RegulationEntry.model_validate(_reg_entry(verdict_status="rewritable"))
    with pytest.raises(ValidationError, match="allowed"):
        D.RegulationEntry.model_validate(_reg_entry(verdict_status="allowed", alternative_expression=["a"]))
    with pytest.raises(ValidationError):
        D.RegulationEntry.model_validate(_reg_entry(verdict_status="banned"))
    with pytest.raises(ValidationError):
        D.RegulationEntry.model_validate(_reg_entry(id="LC-01"))


def test_local_entry_korean_aliases_roundtrip():
    raw = {"id": "LC-90", "항목": "a", "패턴": ["p"], "판정": "irrelevant", "제외하는 맥락": "e", "제외하지 않는 맥락": "k",
           "셀러 문장": "s", "verified_at": "2026-09-28"}
    e = D.LocalEntry.model_validate(raw)
    assert e.item == "a" and e.patterns == ["p"] and e.verdict_status == "irrelevant"
    assert json.loads(D.dump_model(D.LocalDict(dictionary_version="local@2026-09-28.0", source=_src(1), entries=[e])))["entries"][0] == raw
    with pytest.raises(ValidationError):
        D.LocalEntry.model_validate({**raw, "판정": "must_fix"})  # 10월 이후 명칭은 이번 계약 값이 아니다(#7-6)


def _src(n):
    return {"file": "f.xlsx", "sha256": "0" * 64, "sheets": ["s"], "claimed_version": None, "version_status": "unknown",
            "version_note": None, "extracted_at": "2026-09-28", "row_count": n}


def test_dict_rejects_duplicate_ids_and_bad_version_id():
    e = D.LocalEntry.model_validate({"id": "LC-90", "항목": "a", "패턴": ["p"], "판정": "irrelevant", "제외하는 맥락": "e",
                                     "제외하지 않는 맥락": "k", "셀러 문장": "s", "verified_at": "2026-09-28"})
    with pytest.raises(ValidationError, match="중복 id"):
        D.LocalDict(dictionary_version="local@2026-09-28.0", source=_src(2), entries=[e, e])
    with pytest.raises(ValidationError):
        D.LocalDict(dictionary_version="local@v5?/2026-09-22", source=_src(1), entries=[e])  # 불확실성은 식별자에 넣지 않는다(D8)


def test_policy_rules_require_all_four_api_classes():
    raw = json.loads((SYNTH / D.RULES_FILE).read_text(encoding="utf-8"))
    del raw["regulatory_class_map"]["unknown"]
    with pytest.raises(ValidationError, match="unknown"):
        D.PolicyRules.model_validate(raw)
    raw2 = json.loads((SYNTH / D.RULES_FILE).read_text(encoding="utf-8"))
    raw2["overrides"] = [{"keep": "RG-903", "drop": ["RG-902"], "basis": ""}]
    with pytest.raises(ValidationError):
        D.PolicyRules.model_validate(raw2)  # 근거 없는 예외 쌍은 거부


# ---------------------------------------------------------------------------
# 변환 도구 — 임시 xlsx(실제 값 아님)
# ---------------------------------------------------------------------------
REG_HEADER = list(B.REG_COLUMNS) + ["confidence", "source", "version"]
EV_HEADER = list(B.EVIDENCE_COLUMNS)


def _reg_row(**over):
    row = {
        "id": "RG-001", "dict_type": "regulatory", "target_country": "US", "regulatory_class": "cosmetic",
        "source_expression": "표현", "variant_expressions.ko": "표현; 표 현", "variant_expressions.en": "expr; exprs",
        "alternative_expression": "alt", "verdict_status": "regulated", "reason": "사유",
        "evidence_id": "WL-001", "evidence_source_type": "Warning Letter", "evidence_article": "201(g)", "evidence_url": "https://x",
        "verified_at": "2026-09-16", "internal_category": "", "confidence": "high", "source": "담당자 이름이 있는 노트", "version": "3",
    }
    row.update(over)
    return [row[h] for h in REG_HEADER]


def _write_reg_xlsx(path, reg_rows, ev_rows, readme=(("2026-09-15 v2", "x"), ("2026-09-22 v9", "y"))):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = B.README_SHEET
    for r in readme:
        ws.append(list(r))
    ws2 = wb.create_sheet(B.REG_SHEET)
    ws2.append(REG_HEADER)
    for r in reg_rows:
        ws2.append(r)
    ws3 = wb.create_sheet(B.EVIDENCE_SHEET)
    ws3.append(EV_HEADER)
    for r in ev_rows:
        ws3.append(r)
    wb.save(path)


def _ev_row(evidence_id="WL-001", rg_id="RG-001", primary="Y", article="201(g)", url="https://x", quote="quoted"):
    return [evidence_id, rg_id, "Warning Letter", "Co", "2024-07-25", quote, article, url, primary, "2026-09-15"]


def _write_local_xlsx(path, rows, readme=(("현지 부적합 사전 (미국) — 읽는 법",), ("v5 · 대조일 2026-09-22 · 원천: x",))):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = B.README_SHEET
    for r in readme:
        ws.append(list(r))
    ws2 = wb.create_sheet(B.LOCAL_SHEET)
    ws2.append(list(B.LOCAL_COLUMNS) + ["판단 근거", "kr_freq", "kr_corpus"])
    for r in rows:
        ws2.append(r)
    wb.save(path)


def _local_row(rid="LC-01", verdict="needs_fix"):
    return [rid, "항목", "패턴1; 패턴 2", verdict, "제외 맥락", "유지 맥락", "셀러 문장", "2026-09-22", "내부 판단 근거", 41, 66]


def test_build_regulation_extracts_only_allowed_columns_and_evidence(tmp_path):
    x = tmp_path / "r.xlsx"
    _write_reg_xlsx(x, [_reg_row(), _reg_row(id="RG-002", verdict_status="allowed", alternative_expression="", evidence_id="MP-001",
                                              evidence_source_type="시장 관행", evidence_article="", evidence_url="https://mp")],
                    [_ev_row(), _ev_row("WL-009", "RG-001", "N"), _ev_row("MP-001", "RG-002", "Y", "", "https://mp", "market")])
    reg = B.build_regulation(x, date="2026-09-28", seq=1)
    assert reg.dictionary_version == "regulation@2026-09-28.1"
    assert reg.source.claimed_version == "v9 (2026-09-22)" and reg.source.version_status == "as_claimed"
    assert reg.source.sha256 == D.sha256_of(x) and reg.source.row_count == 2
    e = reg.entries[0]
    assert e.variant_ko == ["표현", "표 현"] and e.variant_en == ["expr", "exprs"] and e.alternative_expression == ["alt"]
    assert e.evidence.quote == "quoted" and e.evidence.article == "201(g)"
    dumped = json.loads(D.dump_model(reg))
    for bad in ("confidence", "source", "version"):
        assert bad not in dumped["entries"][0]
    assert "담당자" not in D.dump_model(reg)
    assert reg.entries[1].evidence.article is None and reg.entries[1].alternative_expression == []


@pytest.mark.parametrize(
    "row_over, ev, msg",
    [
        ({"confidence": "medium"}, [_ev_row()], "confidence"),
        ({"verified_at": ""}, [_ev_row()], "verified_at"),
        ({}, [_ev_row(primary="N")], "is_primary"),
        ({}, [_ev_row(url="https://other")], "근거 탭과 다르다"),
        ({"verdict_status": "rewritable", "alternative_expression": ""}, [_ev_row()], "rewritable"),
    ],
)
def test_build_regulation_rejects_invalid_rows(tmp_path, row_over, ev, msg):
    x = tmp_path / "r.xlsx"
    _write_reg_xlsx(x, [_reg_row(**row_over)], ev)
    with pytest.raises(B.BuildError, match=msg):
        B.build_regulation(x, date="2026-09-28", seq=1)


def test_build_regulation_requires_columns(tmp_path):
    import openpyxl
    x = tmp_path / "r.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = B.README_SHEET
    ws = wb.create_sheet(B.REG_SHEET)
    ws.append(["id", "dict_type"])
    ws.append(["RG-001", "regulatory"])
    wb.create_sheet(B.EVIDENCE_SHEET).append(EV_HEADER)
    wb.save(x)
    with pytest.raises(B.BuildError, match="없는 열"):
        B.build_regulation(x, date="2026-09-28", seq=1)


def test_build_local_records_claimed_version_and_mismatch_note(tmp_path):
    x = tmp_path / "l.xlsx"
    _write_local_xlsx(x, [_local_row(), _local_row("LC-02", "irrelevant")])
    loc = B.build_local(x, date="2026-09-28", seq=2, version_status="mismatch_pending", version_note="설명서 v7")
    assert loc.dictionary_version == "local@2026-09-28.2"
    assert loc.source.claimed_version == "v5 (대조일 2026-09-22)"
    assert loc.source.version_status == "mismatch_pending" and loc.source.version_note == "설명서 v7"
    assert loc.entries[0].patterns == ["패턴1", "패턴 2"]
    dumped = json.loads(D.dump_model(loc))["entries"][0]
    assert set(dumped) == set(B.LOCAL_COLUMNS)  # 판단 근거 · kr_freq · kr_corpus 없음
    with pytest.raises(B.BuildError, match="판정"):
        _write_local_xlsx(x, [_local_row(verdict="")])
        B.build_local(x, date="2026-09-28", seq=1, version_status="as_claimed", version_note=None)


def test_build_rules_has_empty_overrides_and_declares_provisional():
    rules = B.build_rules(date="2026-09-28", seq=1)
    assert rules.rules_version == "policy@2026-09-28.1"
    assert rules.overrides == []
    assert "잠정" in rules.basis and "채택" in rules.basis
    statuses = {(r.dict_type, r.verdict_status, r.alternative): (r.emit_verdict_status, r.bucket) for r in rules.verdict_map}
    assert statuses[("regulatory", "allowed", "any")] == (None, "none")
    assert statuses[("regulatory", "rewritable", "required")] == ("regulated", "include")
    assert statuses[("regulatory", "regulated", "absent")] == ("regulated", "exclude")
    assert statuses[("local", "needs_fix", "any")] == ("needs_fix", "exclude")


def test_write_bundle_roundtrip_and_diff(tmp_path):
    rx, lx = tmp_path / "r.xlsx", tmp_path / "l.xlsx"
    _write_reg_xlsx(rx, [_reg_row()], [_ev_row()])
    _write_local_xlsx(lx, [_local_row()])
    old = tmp_path / "old"
    B.write_bundle(old, B.build_regulation(rx, date="2026-09-28", seq=1), B.build_local(lx, date="2026-09-28", seq=1, version_status="unknown", version_note=None), B.build_rules(date="2026-09-28", seq=1))
    B.verify_bundle(old)
    _write_reg_xlsx(rx, [_reg_row(**{"variant_expressions.ko": "표현; 새변형"}), _reg_row(id="RG-003", evidence_id="WL-003")],
                    [_ev_row(), _ev_row("WL-003", "RG-003")])
    _write_local_xlsx(lx, [[*_local_row()[:4], "바뀐 제외 맥락", *_local_row()[5:]]])
    new = tmp_path / "new"
    B.write_bundle(new, B.build_regulation(rx, date="2026-09-28", seq=2), B.build_local(lx, date="2026-09-28", seq=2, version_status="unknown", version_note=None), B.build_rules(date="2026-09-28", seq=2))
    d = B.diff_bundles(old, new)
    assert d["pattern"] == ["RG-001.variant_expressions.ko"]
    assert d["context"] == ["LC-01.제외하는 맥락"]
    assert d["added"] == ["RG-003"] and d["removed"] == [] and d["policy_field"] == []
    assert d["rules"] == ["policy@2026-09-28.1 → policy@2026-09-28.2"]


def test_cli_build_and_verify(tmp_path, capsys):
    rx, lx = tmp_path / "r.xlsx", tmp_path / "l.xlsx"
    _write_reg_xlsx(rx, [_reg_row()], [_ev_row()])
    _write_local_xlsx(lx, [_local_row()])
    out = tmp_path / "out"
    assert B.main(["build", "--regulation", str(rx), "--local", str(lx), "--out", str(out), "--date", "2026-09-28", "--seq", "3"]) == 0
    assert (out / B.SUMS_FILE).exists()
    assert B.main(["verify", "--dir", str(out)]) == 0
    assert B.main(["verify", "--dir", str(tmp_path)]) == 1
    assert "오류" in capsys.readouterr().err
